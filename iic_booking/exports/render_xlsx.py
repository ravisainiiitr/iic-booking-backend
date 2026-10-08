"""Excel workbook: optional Summary sheet (KPIs), one styled sheet per table and a Filters sheet."""

from __future__ import annotations

import io
import re

from . import spec
from .branding import BRAND_HEX
from .branding import PORTAL_LINE
from .values import display_text
from .values import guard_formula
from .values import raw_value
from .values import typed_value

_NUMBER_FORMATS = {
    spec.INTEGER: "#,##0",
    spec.NUMBER: "#,##0.00",
    spec.CURRENCY: '"₹"#,##0.00',
    spec.PERCENT: "0.0%",
    spec.DATE: "DD-MMM-YYYY",
    spec.DATETIME: "DD-MMM-YYYY HH:MM",
}
_FIXED_WIDTHS = {spec.DATE: 13, spec.DATETIME: 18, spec.BOOL: 8}
_HEADER_ROW = 6
_WIDTH_SAMPLE = 500
_MAX_WIDTH = 55
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub(" ", name or "Sheet").strip()[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D0D7E2")
    return {
        "title": Font(bold=True, size=14, color=BRAND_HEX),
        "subtitle": Font(bold=True, size=11, color="1E293B"),
        "meta": Font(italic=True, size=9, color="64748B"),
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill(start_color=BRAND_HEX, end_color=BRAND_HEX, fill_type="solid"),
        "header_align": Alignment(vertical="center", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "total_font": Font(bold=True),
        "total_fill": PatternFill(start_color="E8EEF7", end_color="E8EEF7", fill_type="solid"),
        "wrap": Alignment(vertical="top", wrap_text=True),
        "top": Alignment(vertical="top"),
        "label": Font(bold=True, color="334155"),
        "kpi_value": Font(bold=True, size=12, color=BRAND_HEX),
    }


def _merge_text(ws, row: int, ncols: int, value, font) -> None:
    cell = ws.cell(row=row, column=1, value=guard_formula(value))
    cell.font = font
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)


def _banner(ws, document: spec.Document, heading: str, ncols: int, generated_at: str, styles) -> None:
    _merge_text(ws, 1, ncols, document.title, styles["title"])
    _merge_text(ws, 2, ncols, heading, styles["subtitle"])
    _merge_text(ws, 3, ncols, f"{PORTAL_LINE} · Generated {generated_at} IST", styles["meta"])
    applied = "; ".join(f"{label}: {value}" for label, value in document.filters)
    if applied:
        _merge_text(ws, 4, ncols, f"Filters — {applied}", styles["meta"])


def _write_table(ws, document, table: spec.Table, generated_at: str, styles) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    columns = table.columns
    ncols = max(len(columns), 1)
    count = len(table.rows)
    heading = table.title if len(document.tables) > 1 or not document.subtitle else document.subtitle
    _banner(ws, document, f"{heading} · {count:,} row{'s' if count != 1 else ''}", ncols, generated_at, styles)

    widths = []
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=index, value=guard_formula(column.header))
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_align"]
        cell.border = styles["border"]
        widths.append(max(min(len(column.header), 28), 6))

    totals = {i: 0.0 for i, c in enumerate(columns) if c.total}
    row_no = _HEADER_ROW
    for position, row in enumerate(table.rows):
        row_no += 1
        for index, column in enumerate(columns):
            raw = raw_value(row, column)
            value = typed_value(raw, column.type)
            cell = ws.cell(row=row_no, column=index + 1, value=guard_formula(value))
            cell.border = styles["border"]
            if column.type in _NUMBER_FORMATS and not isinstance(value, str):
                cell.number_format = _NUMBER_FORMATS[column.type]
            if column.type == spec.TEXT:
                cell.alignment = styles["wrap"]
            else:
                cell.alignment = styles["top"]
            if index in totals and isinstance(value, (int, float)):
                totals[index] += value
            if position < _WIDTH_SAMPLE and column.type not in _FIXED_WIDTHS:
                text = display_text(raw, column.type)
                longest = max((len(part) for part in text.split("\n")), default=0)
                widths[index] = max(widths[index], longest + 1)

    if not table.rows:
        row_no += 1
        _merge_text(ws, row_no, ncols, table.empty_message, styles["meta"])
    elif totals:
        row_no += 1
        for index, column in enumerate(columns):
            cell = ws.cell(row=row_no, column=index + 1)
            if index == 0:
                cell.value = "Total"
            elif index in totals:
                cell.value = totals[index]
                cell.number_format = _NUMBER_FORMATS.get(column.type, "General")
            cell.font = styles["total_font"]
            cell.fill = styles["total_fill"]
            cell.border = styles["border"]
    if table.note:
        row_no += 2
        _merge_text(ws, row_no, ncols, table.note, styles["meta"])

    for index, column in enumerate(columns, start=1):
        width = _FIXED_WIDTHS.get(column.type) or min(max(widths[index - 1] + 1, 8), _MAX_WIDTH)
        ws.column_dimensions[get_column_letter(index)].width = width

    if columns:
        last = get_column_letter(len(columns))
        ws.freeze_panes = ws.cell(row=_HEADER_ROW + 1, column=1)
        if table.rows:
            ws.auto_filter.ref = f"A{_HEADER_ROW}:{last}{_HEADER_ROW + count}"
        ws.print_title_rows = f"{_HEADER_ROW}:{_HEADER_ROW}"
    ws.page_setup.orientation = "landscape" if len(columns) > 6 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _write_summary(ws, document: spec.Document, generated_at: str, styles) -> None:
    _banner(ws, document, document.subtitle or "Summary", 3, generated_at, styles)
    row = _HEADER_ROW
    for index, header in enumerate(("Metric", "Value", "Notes"), start=1):
        cell = ws.cell(row=row, column=index, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.border = styles["border"]
    for kpi in document.kpis:
        row += 1
        label = ws.cell(row=row, column=1, value=guard_formula(kpi.label))
        label.font = styles["label"]
        value = typed_value(kpi.value, kpi.type)
        cell = ws.cell(row=row, column=2, value=guard_formula(value))
        cell.font = styles["kpi_value"]
        if kpi.type in _NUMBER_FORMATS and not isinstance(value, str):
            cell.number_format = _NUMBER_FORMATS[kpi.type]
        ws.cell(row=row, column=3, value=guard_formula(kpi.hint or None))
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).border = styles["border"]
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 50


def _write_filters(ws, document: spec.Document, generated_at: str, styles) -> None:
    ws.cell(row=1, column=1, value=guard_formula(document.title)).font = styles["title"]
    ws.cell(row=2, column=1, value=PORTAL_LINE).font = styles["meta"]
    rows = [("Generated at (IST)", generated_at)]
    if document.generated_by:
        rows.append(("Generated by (role)", document.generated_by))
    for table in document.tables:
        label = f"Rows — {table.title}" if len(document.tables) > 1 else "Rows"
        rows.append((label, f"{len(table.rows):,}"))
    rows.extend(document.filters or [("Filters", "None (all records you can see)")])
    for offset, (label, value) in enumerate(rows, start=4):
        ws.cell(row=offset, column=1, value=guard_formula(label)).font = styles["label"]
        ws.cell(row=offset, column=2, value=guard_formula(str(value)))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70


def render_xlsx(document: spec.Document, *, generated_at: str) -> bytes:
    from openpyxl import Workbook

    styles = _styles()
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    if document.kpis:
        _write_summary(wb.create_sheet(_sheet_title("Summary", used)), document, generated_at, styles)
    for table in document.tables:
        ws = wb.create_sheet(_sheet_title(table.sheet_name or table.title, used))
        _write_table(ws, document, table, generated_at, styles)
    _write_filters(wb.create_sheet(_sheet_title("Filters", used)), document, generated_at, styles)
    wb.properties.title = document.title
    wb.properties.creator = PORTAL_LINE
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
