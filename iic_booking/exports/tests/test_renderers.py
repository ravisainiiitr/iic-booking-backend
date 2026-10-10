"""Shared export module: CSV / Excel / PDF rendering, value types, formula guard, caps and filenames."""

import csv
import io
import re
from datetime import date
from datetime import datetime
from datetime import timezone as dt_timezone

import pytest
from openpyxl import load_workbook

from iic_booking.exports import spec
from iic_booking.exports.http import EXPORT_ROW_LIMIT
from iic_booking.exports.http import PDF_ROW_LIMIT
from iic_booking.exports.http import ExportTooLarge
from iic_booking.exports.http import export_filename
from iic_booking.exports.http import export_response
from iic_booking.exports.render_csv import render_csv
from iic_booking.exports.render_pdf import render_pdf
from iic_booking.exports.render_xlsx import render_xlsx
from iic_booking.exports.values import display_text
from iic_booking.exports.values import guard_formula
from iic_booking.exports.values import indian_grouping

C = spec.Column

COLUMNS = [
    C("name", "Name", width=1.5),
    C("count", "Count", spec.INTEGER, total=True),
    C("amount", "Amount (₹)", spec.CURRENCY, total=True),
    C("share", "Share", spec.PERCENT),
    C("day", "Day", spec.DATE),
    C("at", "At (IST)", spec.DATETIME),
    C("ok", "OK", spec.BOOL),
    C("nested", "Nested", value=lambda r: (r.get("meta") or {}).get("label")),
]
ROWS = [
    {"name": "Alice", "count": 3, "amount": "1500.50", "share": 0.25, "day": "2026-10-01",
     "at": "2026-10-01T04:30:00+00:00", "ok": True, "meta": {"label": "first"}},
    {"name": '=HYPERLINK("http://x")', "count": 2, "amount": -20, "share": 0.5, "day": date(2026, 10, 2),
     "at": datetime(2026, 10, 2, 10, 0, tzinfo=dt_timezone.utc), "ok": False, "meta": None},
    {"name": "हिंदी नाम", "count": None, "amount": None, "share": None, "day": None, "at": None, "ok": None},
]


def _doc(rows=ROWS, **kw):
    table = spec.Table("main", "People", list(COLUMNS), [dict(r) for r in rows])
    defaults = dict(title="Test Report", slug="test-report", tables=[table], subtitle="October",
                    filters=[("Status", "Active"), ("From", "01 Oct 2026")],
                    kpis=[spec.Kpi("Total", 1520.5, spec.CURRENCY, "hint"), spec.Kpi("Rows", 3, spec.INTEGER)],
                    generated_by="Admin")
    defaults.update(kw)
    return spec.Document(**defaults)


def test_guard_formula_prefixes_only_text():
    assert guard_formula("=1+1") == "'=1+1"
    assert guard_formula("+91 999") == "'+91 999"
    assert guard_formula("-x") == "'-x"
    assert guard_formula("@SUM(A1)") == "'@SUM(A1)"
    assert guard_formula("plain") == "plain"
    assert guard_formula(-5) == -5


def test_display_text_types():
    assert display_text("1234567.5", spec.CURRENCY) == "₹12,34,567.50"
    assert display_text("1234567.5", spec.CURRENCY, machine=True) == "1234567.50"
    assert display_text(0.256, spec.PERCENT) == "25.6%"
    assert display_text("2026-10-01T04:30:00+00:00", spec.DATETIME) == "01 Oct 2026, 10:00"
    assert display_text("2026-10-01T04:30:00+00:00", spec.DATETIME, machine=True) == "2026-10-01 10:00"
    assert display_text("2026-10-01", spec.DATE) == "01 Oct 2026"
    assert display_text(True, spec.BOOL) == "Yes"
    assert display_text(12000, spec.INTEGER) == "12,000"
    assert indian_grouping(-123456789.0) == "-12,34,56,789.00"


def test_csv_bom_escaping_guard_and_types():
    raw = render_csv(_doc())
    assert raw.startswith("\ufeff".encode())
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == ["S.No.", *[c.header for c in COLUMNS]]
    assert rows[1] == ["1", "Alice", "3", "1500.50", "25.0%", "2026-10-01", "2026-10-01 10:00", "Yes", "first"]
    assert rows[2][:2] == ["2", "'=HYPERLINK(\"http://x\")"]
    assert rows[2][3] == "-20.00"
    assert rows[3][1] == "हिंदी नाम"
    assert len(rows) == 4


def test_csv_multiple_tables_are_sections():
    doc = _doc()
    doc.tables.append(spec.Table("second", "Second table", [C("x", "X")], [{"x": "1"}]))
    rows = list(csv.reader(io.StringIO(render_csv(doc).decode("utf-8-sig"))))
    assert rows[0] == ["People"]
    assert [] in rows
    assert rows[-3:] == [["Second table"], ["S.No.", "X"], ["1", "1"]]


def test_serial_column_is_optional_and_not_doubled():
    from iic_booking.exports.reports.common import SNO
    from iic_booking.exports.reports.common import numbered

    breakdown = spec.Table("status", "By status", [C("x", "X")], [{"x": "a"}], serial=False)
    own = spec.Table("own", "Own", [SNO, C("x", "X")], numbered([{"x": "a"}, {"x": "b"}]))
    doc = spec.Document(title="T", slug="t", tables=[breakdown, own])
    rows = list(csv.reader(io.StringIO(render_csv(doc).decode("utf-8-sig"))))
    assert rows[1:3] == [["X"], ["a"]]
    assert rows[-3:] == [["S.No.", "X"], ["1", "a"], ["2", "b"]]


def test_xlsx_sheets_styles_formats_and_guard():
    wb = load_workbook(io.BytesIO(render_xlsx(_doc(), generated_at="08 Oct 2026, 10:00")))
    assert wb.sheetnames == ["Summary", "People", "Filters"]
    ws = wb["People"]
    assert ws["A1"].value == "Test Report"
    assert [ws.cell(row=6, column=i).value for i in range(1, 10)] == ["S.No.", *[c.header for c in COLUMNS]]
    assert ws["A6"].fill.start_color.rgb.endswith("153F79")
    assert ws["B6"].alignment.horizontal == "center" and ws["B6"].alignment.vertical == "center"
    assert ws.freeze_panes == "A7"
    assert ws.auto_filter.ref == "A6:I9"
    assert ws["A7"].value == 1 and ws["A9"].value == 3
    assert ws["B7"].alignment.horizontal == "center" and ws["B7"].alignment.vertical == "center"
    assert ws["C7"].value == 3 and ws["C7"].number_format == "#,##0"
    assert ws["D7"].value == 1500.5 and "₹" in ws["D7"].number_format
    assert ws["E7"].number_format == "0.0%"
    assert isinstance(ws["F7"].value, datetime) and ws["F7"].number_format == "DD-MMM-YYYY"
    assert ws["G7"].value == datetime(2026, 10, 1, 10, 0)
    assert ws["B8"].value == "'=HYPERLINK(\"http://x\")"
    assert ws["A10"].value == "Total" and ws["C10"].value == 5 and ws["D10"].value == 1480.5
    summary = wb["Summary"]
    assert summary["A7"].value == "Total" and summary["B7"].value == 1520.5
    filters = {r[0]: r[1] for r in wb["Filters"].iter_rows(min_row=4, values_only=True) if r[0]}
    assert filters["Generated at (IST)"] == "08 Oct 2026, 10:00"
    assert filters["Generated by (role)"] == "Admin"
    assert filters["Status"] == "Active"


def _linked_doc():
    columns = [
        C("booking_id", "Booking ID", link=lambda r: f"/booking-management?expand={r['pk']}"),
        C("note", "Note", align="left"),
    ]
    rows = [{"booking_id": "IIC-7", "pk": 7, "note": "Long free text"}]
    return spec.Document(title="Links", slug="links", tables=[spec.Table("b", "Bookings", columns, rows)])


def test_xlsx_links_and_left_aligned_text(settings):
    settings.FRONTEND_URL = "https://equip.iitr.ac.in"
    ws = load_workbook(io.BytesIO(render_xlsx(_linked_doc(), generated_at="x")))["Bookings"]
    cell = ws["B7"]
    assert cell.value == "IIC-7"
    assert cell.hyperlink.target == "https://equip.iitr.ac.in/booking-management?expand=7"
    assert cell.font.underline == "single"
    assert ws["C7"].alignment.horizontal == "left"


def test_pdf_links_booking_ids(settings):
    settings.FRONTEND_URL = "https://equip.iitr.ac.in"
    raw = render_pdf(_linked_doc(), generated_at="x")
    assert b"https://equip.iitr.ac.in/booking-management?expand=7" in raw


def test_links_are_omitted_without_frontend_url(settings):
    settings.FRONTEND_URL = ""
    ws = load_workbook(io.BytesIO(render_xlsx(_linked_doc(), generated_at="x")))["Bookings"]
    assert ws["B7"].hyperlink is None


def test_booking_detail_paths_by_role():
    from iic_booking.equipment.booking_links import booking_detail_path

    assert booking_detail_path(pk=12, display_id="IIC-12", staff=True) == "/booking-management?expand=12"
    assert booking_detail_path(pk=None, display_id="IIC-12", staff=True) == ""
    assert booking_detail_path(pk=12, display_id="IIC 12/A", staff=False) == "/my-bookings?booking=IIC%2012%2FA"
    assert booking_detail_path(pk=12, display_id="", staff=False) == ""


def test_xlsx_empty_table_has_message():
    wb = load_workbook(io.BytesIO(render_xlsx(_doc(rows=[], kpis=[]), generated_at="x")))
    assert wb.sheetnames == ["People", "Filters"]
    assert wb["People"]["A7"].value == "No records match these filters."


def _pdf_page_count(raw: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", raw))


def test_pdf_has_unicode_fonts_and_pages():
    raw = render_pdf(_doc(), generated_at="08 Oct 2026, 10:00")
    assert raw.startswith(b"%PDF")
    assert b"IICNotoSans" in raw or b"NotoSans" in raw
    assert b"NotoSansDevanagari" in raw or b"IICNotoDeva" in raw
    assert _pdf_page_count(raw) == 1


def test_pdf_many_rows_repeat_pages_and_landscape():
    rows = [{"name": f"User {i}", "count": i, "amount": i * 10} for i in range(400)]
    raw = render_pdf(_doc(rows=rows), generated_at="x")
    assert _pdf_page_count(raw) > 3
    assert b"/MediaBox [ 0 0 841.8898 595.2756 ]" in raw


def test_pdf_long_titled_table_starts_on_first_page():
    from pypdf import PdfReader

    columns = [spec.Column("name", "Name"), spec.Column("count", "Count", spec.INTEGER)]
    rows = [{"name": f"Row {i}", "count": i} for i in range(80)]
    document = spec.Document(title="Log", slug="log", tables=[spec.Table("log", "All entries", columns, rows)])
    first_page = PdfReader(io.BytesIO(render_pdf(document, generated_at="x"))).pages[0].extract_text()
    assert "Row 0" in first_page


def test_export_response_caps_and_headers():
    response = export_response(_doc(), "csv")
    assert response["X-Export-Row-Count"] == "3"
    assert re.fullmatch(r'attachment; filename="test-report_\d{4}-\d{2}-\d{2}_\d{4}\.csv"',
                        response["Content-Disposition"])
    big = [{"name": "x"}] * (PDF_ROW_LIMIT + 1)
    with pytest.raises(ExportTooLarge, match="PDF exports are limited"):
        export_response(_doc(rows=big), "pdf")
    assert export_response(_doc(rows=big), "csv").status_code == 200
    with pytest.raises(ExportTooLarge, match="limited to 10,000"):
        export_response(_doc(rows=[{"name": "x"}] * (EXPORT_ROW_LIMIT + 1)), "xlsx")


def test_filename_uses_slug_and_ist_time():
    moment = datetime(2026, 10, 8, 7, 5)
    assert export_filename("booking-attempt-log", "xlsx", moment) == "booking-attempt-log_2026-10-08_0705.xlsx"


def test_unknown_column_type_rejected():
    with pytest.raises(ValueError):
        C("x", "X", "money")
