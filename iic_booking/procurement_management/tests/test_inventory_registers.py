"""Asset registers, bulk import, verification, disposal, roles (Accounts / Lab In Charge), OC Stores line edits,
equipment-linked inventory, maintenance, the back-to-functional hook, purchase mode and bills for Accounts."""

import io
import uuid
from datetime import timedelta
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from iic_booking.procurement_management import config_service
from iic_booking.procurement_management import constants as c
from iic_booking.procurement_management.models import (
    ApprovalAction,
    Asset,
    AssetRegister,
    Invoice,
    Item,
    MaintenanceRecord,
    ProcurementRecord,
    PurchaseRequest,
    StockBalance,
)
from iic_booking.users.models.user_type import UserType

from .conftest import API, act, category, client_for, line, make_equipment, make_user, new_request

pytestmark = pytest.mark.django_db
R = c.ModuleRole
TODAY = timezone.localdate


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def make_register(world, code="MAJ-1", rtype="MAJOR", user=None, expect=201, **extra):
    res = client_for(user or world.stores).post(
        f"{API}/registers/", {"department_id": world.dept.pk, "code": code, "register_type": rtype, "name": f"{code} book", **extra},
        format="json",
    )
    assert res.status_code == expect, res.json()
    return res.json()


def add_asset(world, user=None, expect=201, **extra):
    body = {"equipment_id": world.equipment.pk, "category_id": category(world.dept, "MAJOR_ASSET").pk,
            "description": "FE-SEM column", "cost": "2500000.00", **extra}
    res = client_for(user or world.stores).post(f"{API}/assets/", body, format="json")
    assert res.status_code == expect, res.json()
    return res.json()


def make_item(world, cat="CONSUMABLE", name="Vacuum pump oil"):
    return Item.objects.create(
        department=world.dept, code=f"ITM-{uuid.uuid4().hex[:6]}", name=name, category=category(world.dept, cat), uom="L"
    )


def stock_in(world, item, qty, price="500.00"):
    res = client_for(world.stores).post(
        f"{API}/stock/transactions/",
        {"department_id": world.dept.pk, "item_id": item.pk, "tx_type": "RECEIPT", "quantity": qty, "unit_cost": price},
        format="json",
    )
    assert res.status_code == 201, res.json()


def central(world, item):
    row = StockBalance.objects.filter(department=world.dept, item=item, laboratory__isnull=True).first()
    return row.quantity if row else Decimal("0")


def csv_upload(rows, name="registers.csv"):
    from iic_booking.procurement_management.asset_import import COLUMNS

    headers = [h for h, _, _, _ in COLUMNS]
    keys = [k for _, k, _, _ in COLUMNS]
    out = io.StringIO()
    out.write(",".join(headers) + "\n")
    for r in rows:
        out.write(",".join(str(r.get(k, "")) for k in keys) + "\n")
    return SimpleUploadedFile(name, out.getvalue().encode("utf-8"), content_type="text/csv")


# ---------------------------------------------------------------------------
# Registers and register entries
# ---------------------------------------------------------------------------
class TestRegisters:
    def test_register_crud_and_permissions(self, world):
        reg = make_register(world, code="maj 1", total_pages=200)
        assert reg["code"] == "MAJ-1" and reg["register_type"] == "MAJOR" and reg["entry_count"] == 0
        assert make_register(world, code="MAJ-1", expect=400)["code"] == "duplicate_register"
        make_register(world, code="MIN-1", rtype="MINOR", user=world.operator, expect=403)
        res = client_for(world.stores).patch(f"{API}/registers/{reg['id']}/", {"volume": "II"}, format="json")
        assert res.status_code == 200 and res.json()["volume"] == "II"
        assert client_for(world.other_oic).get(f"{API}/registers/{reg['id']}/").status_code == 404

    def test_entry_unique_tag_search_and_lookup(self, world):
        reg = make_register(world, total_pages=100)
        a = add_asset(world, register_id=reg["id"], register_page=12, register_serial="3")["results"][0]
        assert a["register_ref"] == "MAJ-1 / p.12 / s.3"
        assert a["asset_tag"] == f"{world.dept.code}/MAJ/{a['id']:06d}"
        dup = add_asset(world, expect=400, register_id=reg["id"], register_page=12, register_serial="3")
        assert dup["code"] == "duplicate_entry" and dup["clashes"][0]["asset_id"] == a["id"]
        assert add_asset(world, expect=400, register_id=reg["id"], register_page=101, register_serial="1")["code"] == "invalid"
        assert add_asset(world, expect=400, register_page=4)["field"] == "register_id"

        found = client_for(world.stores).get(f"{API}/assets/", {"q": "MAJ-1/12/3"}).json()
        assert [x["id"] for x in found["results"]] == [a["id"]]
        found = client_for(world.stores).get(f"{API}/assets/", {"q": "maj-1 p12 s3"}).json()
        assert [x["id"] for x in found["results"]] == [a["id"]]
        hit = client_for(world.operator).get(f"{API}/assets/lookup/", {"tag": a["asset_tag"]})
        assert hit.status_code == 200 and hit.json()["id"] == a["id"] and hit.json()["can_verify"] is True
        hit = client_for(world.stores).get(f"{API}/assets/lookup/", {"register_id": reg["id"], "page_no": 12, "serial": "3"})
        assert hit.json()["id"] == a["id"]
        assert client_for(world.other_oic).get(f"{API}/assets/lookup/", {"tag": a["asset_tag"]}).status_code == 404

    def test_batch_increments_serials_and_edit_checks_duplicates(self, world):
        reg = make_register(world)
        out = add_asset(world, count=3, register_id=reg["id"], register_page=5, register_serial="7")["results"]
        assert [x["register_serial"] for x in out] == ["7", "8", "9"]
        res = client_for(world.stores).patch(f"{API}/assets/{out[0]['id']}/", {"register_serial": "8"}, format="json")
        assert res.status_code == 400 and res.json()["code"] == "duplicate_entry"
        res = client_for(world.stores).patch(
            f"{API}/assets/{out[0]['id']}/", {"register_serial": "10", "supplier_name": "Zeiss", "condition": "GOOD"}, format="json"
        )
        assert res.status_code == 200 and res.json()["register_serial"] == "10" and res.json()["condition"] == "GOOD"

    def test_accessory_parent_and_detail(self, world):
        main = add_asset(world)["results"][0]
        acc = add_asset(world, description="Detector", parent_id=main["id"])["results"][0]
        assert acc["parent"]["id"] == main["id"]
        detail = client_for(world.stores).get(f"{API}/assets/{main['id']}/").json()
        assert [x["id"] for x in detail["accessories"]] == [acc["id"]]
        assert add_asset(world, expect=400, parent_id=acc["id"])["code"] == "invalid_parent"

    def test_register_print_and_labels(self, world):
        reg = make_register(world)
        a = add_asset(world, register_id=reg["id"], register_page=1, register_serial="1")["results"][0]
        res = client_for(world.stores).get(f"{API}/registers/{reg['id']}/entries/", {"export": "pdf"})
        assert res.status_code == 200 and res["Content-Type"] == "application/pdf" and res.content[:4] == b"%PDF"
        res = client_for(world.stores).get(f"{API}/registers/{reg['id']}/entries/", {"export": "xlsx"})
        assert res.status_code == 200
        res = client_for(world.stores).post(f"{API}/assets/labels/", {"asset_ids": [a["id"]]}, format="json")
        assert res.status_code == 200 and res.content[:4] == b"%PDF"
        assert client_for(world.stores).post(f"{API}/assets/labels/", {}, format="json").status_code == 400

    def test_capital_items_from_procurement_are_auto_placed(self, world):
        from iic_booking.procurement_management.registers import entry_values

        reg_row = AssetRegister.objects.create(department=world.dept, register_type="MAJOR", code="MAJ-9", name="x")
        add_asset(world, register_id=reg_row.pk, register_page=4, register_serial="6")
        vals = entry_values({}, world.dept.pk, 2, category=category(world.dept, "MAJOR_ASSET"), auto_place=True)
        assert [(v["register"].pk, v["register_page"], v["register_serial"]) for v in vals] == [(reg_row.pk, 4, "7"), (reg_row.pk, 4, "8")]


# ---------------------------------------------------------------------------
# Bulk import
# ---------------------------------------------------------------------------
class TestImport:
    def test_template_download(self, world):
        res = client_for(world.stores).get(f"{API}/assets/import/template/")
        assert res.status_code == 200 and res["Content-Type"].startswith("application/vnd.openxmlformats")
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(res.content))
        assert wb["Assets"]["A1"].value == "Register Code" and "Instructions" in wb.sheetnames
        assert client_for(world.stores).get(f"{API}/assets/import/template/", {"type": "csv"}).status_code == 200

    def test_preview_flags_errors_and_duplicates_then_commit(self, world):
        reg = make_register(world)
        existing = add_asset(world, register_id=reg["id"], register_page=2, register_serial="1")["results"][0]
        rows = [
            {"register_code": "MAJ-1", "register_type": "MAJOR", "page": "2", "serial": "2", "description": "Rotary pump",
             "equipment_code": world.equipment.code, "cost": "45000", "asset_tag": "OLD-77", "entry_date": "14-08-2019"},
            {"register_code": "MAJ-1", "register_type": "MAJOR", "page": "2", "serial": "1", "description": "Dup of DB"},
            {"register_code": "MAJ-1", "register_type": "MAJOR", "page": "2", "serial": "2", "description": "Dup in file"},
            {"register_code": "MIN-4", "register_type": "MINOR", "page": "1", "serial": "1", "description": "Desk lamp",
             "parent_tag": "OLD-77", "condition": "FAIR"},
            {"register_code": "MAJ-1", "register_type": "MAJOR", "page": "x", "serial": "9", "description": "Bad page",
             "equipment_code": "NOPE"},
        ]
        cl = client_for(world.stores)
        prev = cl.post(f"{API}/assets/import/preview/", {"department_id": world.dept.pk, "file": csv_upload(rows)}, format="multipart")
        assert prev.status_code == 200, prev.json()
        body = prev.json()
        status = {r["row"]: r["status"] for r in body["rows"]}
        assert status == {2: "OK", 3: "DUPLICATE", 4: "DUPLICATE", 5: "ERROR", 6: "ERROR"}
        assert body["rows"][1]["duplicate_of"]["asset_id"] == existing["id"]
        assert "MIN-4" in " ".join(body["rows"][3]["errors"])

        bad = cl.post(f"{API}/assets/import/commit/", {"department_id": world.dept.pk, "file": csv_upload(rows)}, format="multipart")
        assert bad.status_code == 400 and bad.json()["code"] == "import_errors"
        assert Asset.objects.filter(department=world.dept).count() == 1

        done = cl.post(
            f"{API}/assets/import/commit/",
            {"department_id": world.dept.pk, "file": csv_upload(rows), "skip_errors": "true", "create_registers": "true"},
            format="multipart",
        )
        assert done.status_code == 201, done.json()
        assert done.json()["created"] == 2 and done.json()["skipped"] == 3 and done.json()["registers_created"] == ["MIN-4"]
        pump = Asset.objects.get(asset_tag="OLD-77")
        assert (pump.register.code, pump.register_page, pump.register_serial) == ("MAJ-1", 2, "2")
        assert pump.equipment_id == world.equipment.pk and pump.register_entry_date.isoformat() == "2019-08-14"
        lamp = Asset.objects.get(description="Desk lamp")
        assert lamp.parent_id == pump.pk and lamp.register.register_type == "MINOR" and lamp.asset_tag
        assert lamp.equipment_id == world.equipment.pk

    def test_import_requires_asset_permission(self, world):
        res = client_for(world.operator).post(
            f"{API}/assets/import/preview/", {"department_id": world.dept.pk, "file": csv_upload([])}, format="multipart"
        )
        assert res.status_code == 403


# ---------------------------------------------------------------------------
# Physical verification and disposal
# ---------------------------------------------------------------------------
class TestVerificationAndDisposal:
    def test_campaign_verify_and_stats(self, world):
        reg = make_register(world)
        a = add_asset(world, register_id=reg["id"], register_page=1, register_serial="1")["results"][0]
        b = add_asset(world, register_id=reg["id"], register_page=1, register_serial="2")["results"][0]
        cmp = client_for(world.stores).post(f"{API}/verification/campaigns/", {"department_id": world.dept.pk, "register_id": reg["id"]}, format="json")
        assert cmp.status_code == 201, cmp.json()
        assert cmp.json()["stats"] == {"total": 2, "verified": 0, "pending": 2, "by_result": {}}
        cid = cmp.json()["id"]
        res = client_for(world.operator).post(
            f"{API}/assets/{a['id']}/verifications/", {"campaign_id": cid, "result": "FOUND", "method": "SCAN", "condition": "GOOD"}, format="json"
        )
        assert res.status_code == 201, res.json()
        res = client_for(world.stores).post(f"{API}/assets/{b['id']}/verifications/", {"campaign_id": cid, "result": "NOT_FOUND"}, format="json")
        assert res.status_code == 400  # remarks needed for a discrepancy
        res = client_for(world.stores).post(
            f"{API}/assets/{b['id']}/verifications/", {"campaign_id": cid, "result": "NOT_FOUND", "remarks": "Not in room 112"}, format="json"
        )
        assert res.status_code == 201
        assert client_for(world.other_oic).post(f"{API}/assets/{a['id']}/verifications/", {"result": "FOUND"}, format="json").status_code == 404
        stats = client_for(world.stores).get(f"{API}/verification/campaigns/{cid}/").json()["stats"]
        assert stats == {"total": 2, "verified": 2, "pending": 0, "by_result": {"FOUND": 1, "NOT_FOUND": 1}}
        asset = Asset.objects.get(pk=a["id"])
        assert asset.last_verified_on == TODAY() and asset.last_verification_result == "FOUND" and asset.condition == "GOOD"
        pending = client_for(world.stores).get(f"{API}/verification/campaigns/{cid}/assets/").json()
        assert pending["count"] == 0
        closed = client_for(world.stores).post(f"{API}/verification/campaigns/{cid}/close/", {}, format="json")
        assert closed.json()["status"] == "CLOSED"
        assert client_for(world.stores).post(
            f"{API}/assets/{a['id']}/verifications/", {"campaign_id": cid, "result": "FOUND"}, format="json"
        ).status_code == 400

    def test_condemn_then_dispose(self, world):
        a = add_asset(world)["results"][0]
        cl = client_for(world.stores)
        url = f"{API}/assets/{a['id']}/dispose/"
        assert cl.post(url, {"action": "DISPOSE", "sanction_reference": "S/1", "remarks": "x"}, format="json").json()["code"] == "invalid_transition"
        assert cl.post(url, {"action": "CONDEMN", "remarks": "Beyond repair"}, format="json").status_code == 400
        res = cl.post(url, {"action": "CONDEMN", "board_reference": "Condemnation board 2026/4", "remarks": "Beyond repair"}, format="json")
        assert res.status_code == 201 and res.json()["asset"]["status"] == "CONDEMNED"
        res = cl.post(url, {"action": "DISPOSE", "mode": "AUCTION", "sanction_reference": "HOD/2026/9", "realised_value": "1200",
                            "remarks": "Auctioned"}, format="json")
        assert res.status_code == 201 and res.json()["asset"]["status"] == "DISPOSED"
        assert len(res.json()["asset"]["disposals"]) == 2
        assert client_for(world.operator).post(url, {"action": "CONDEMN"}, format="json").status_code == 403


# ---------------------------------------------------------------------------
# Roles: Accounts In Charge, Lab In Charge
# ---------------------------------------------------------------------------
class TestRoles:
    def test_accounts_role_permissions_and_budget_stage(self, world):
        accounts = make_user(user_type=UserType.FINANCE, name="Accounts", department=world.dept)
        config_service.assign_role(world.admin, world.dept, accounts, R.ACCOUNTS)
        from iic_booking.procurement_management.access import scope_for

        perms = scope_for(accounts).permissions(world.dept.pk)
        assert {c.OfficePermission.PAYMENTS, c.OfficePermission.INVOICES, c.OfficePermission.BUDGET} <= perms
        assert c.OfficePermission.ASSETS not in perms

        config_service.update_config(world.admin, world.dept, {"accounts_budget_check": True})
        r = new_request(world.operator, equipment=world.equipment, submit=True)
        assert r.approval_route == ["OIC", "STORES", "ACCOUNTS"] and r.stage_entered_at is not None
        act(world.oic, r, "approve")
        act(world.stores, r, "approve")
        assert r.status == c.RequestStatus.PENDING_ACCOUNTS
        inbox = client_for(accounts).get(f"{API}/approvals/").json()
        assert [x["id"] for x in inbox["results"]] == [r.pk]
        act(world.office, r, "approve", expect=403)
        act(accounts, r, "approve", comments="Budget available under Non-Plan 2026-27")
        assert r.status == c.RequestStatus.APPROVED
        dash = client_for(accounts).get(f"{API}/dashboard/", {"department_id": world.dept.pk}).json()
        assert dash["accounts"]["budget_checks_pending"] == 0

    def test_lab_incharge_scope(self, world):
        lic = make_user(user_type=UserType.MANAGER, name="Lab In Charge", department=world.dept)
        res = client_for(world.admin).post(
            f"{API}/config/{world.dept.pk}/roles/",
            {"user_id": lic.pk, "role": R.LAB_INCHARGE, "equipment_ids": [world.equipment2.pk]}, format="json",
        )
        assert res.status_code == 201 and res.json()["equipment_ids"] == [world.equipment2.pk]
        bad = client_for(world.admin).post(
            f"{API}/config/{world.dept.pk}/roles/",
            {"user_id": lic.pk, "role": R.LAB_INCHARGE, "equipment_ids": [world.other_equipment.pk]}, format="json",
        )
        assert bad.status_code == 400 and bad.json()["code"] == "invalid_equipment"
        r = new_request(lic, equipment=world.equipment2)
        assert r.raised_as_role == R.LAB_INCHARGE
        new_request(lic, equipment=world.equipment, expect=403)
        new_request(lic, expect=400)  # Lab In Charge must pick the equipment
        r = new_request(world.operator2, equipment=world.equipment2, submit=True)
        assert r.pk in [x["id"] for x in client_for(lic).get(f"{API}/requests/").json()["results"]]


# ---------------------------------------------------------------------------
# OC Stores line edits
# ---------------------------------------------------------------------------
class TestStoresEdit:
    def test_edit_add_remove_and_history(self, world):
        item = make_item(world)
        sub = make_item(world, name="Pump oil (equivalent)")
        stock_in(world, item, "10")
        r = new_request(
            world.operator, equipment=world.equipment, submit=True,
            lines=[{"item_id": item.pk, "quantity": "4", "estimated_unit_price": "500"}, line(description="Gloves", price="200", qty="2")],
        )
        act(world.oic, r, "approve")
        lines = list(r.lines.order_by("id"))
        stock = client_for(world.stores).get(f"{API}/requests/{r.pk}/stock-check/").json()["results"]
        assert stock[0]["suggested"] == "STOCK" and stock[0]["on_hand"] == "10.000" and stock[1]["suggested"] == ""

        url = f"{API}/requests/{r.pk}/stores-edit/"
        body = {"lines": [
            {"id": lines[0].pk, "quantity": "3", "fulfilment": "STOCK", "store_note": "3 is enough for a refill"},
            {"id": lines[1].pk, "item_id": sub.pk, "estimated_unit_price": "300", "fulfilment": "PROCURE"},
            {"item_id": item.pk, "quantity": "1", "estimated_unit_price": "500", "fulfilment": "STOCK"},
        ], "comments": "Adjusted after checking the store"}
        assert client_for(world.operator).post(url, body, format="json").status_code == 403
        res = client_for(world.stores).post(url, body, format="json")
        assert res.status_code == 200, res.json()
        out = res.json()
        assert out["estimated_total"] == "2600.00"
        first, second, added = out["lines"]
        assert first["store_original"]["quantity"] == "4.000" and first["fulfilment"] == "STOCK"
        assert second["item_id"] == sub.pk and second["description"] == sub.name and second["store_original"]["description"] == "Gloves"
        assert added["added_by_stores"] is True and added["store_original"] == {}
        hist = ApprovalAction.objects.filter(purchase_request=r, action=c.ApprovalActionType.STORES_EDIT).get()
        assert "Estimated total" in hist.comments
        res = client_for(world.stores).post(url, {"lines": [{"id": added["id"], "remove": True}]}, format="json")
        assert len(res.json()["lines"]) == 2

    def test_by_lines_decision_issues_stock_lines(self, world):
        item = make_item(world)
        stock_in(world, item, "10")
        r = new_request(
            world.operator, equipment=world.equipment, submit=True,
            lines=[{"item_id": item.pk, "quantity": "4", "estimated_unit_price": "500"}, line(description="Special glassware")],
        )
        act(world.oic, r, "approve")
        a, b = r.lines.order_by("id")
        client_for(world.stores).post(
            f"{API}/requests/{r.pk}/stores-edit/",
            {"lines": [{"id": a.pk, "fulfilment": "STOCK"}, {"id": b.pk, "fulfilment": "PROCURE"}]}, format="json",
        )
        act(world.stores, r, "stores-review", decision="BY_LINES")
        assert central(world, item) == Decimal("6.000")
        a.refresh_from_db()
        assert a.issued_quantity == Decimal("4.000") and r.status != c.RequestStatus.PENDING_STORES

    def test_raising_total_adds_hod_stage(self, world):
        r = new_request(world.operator, equipment=world.equipment, submit=True, lines=[line(price="1000")])
        act(world.oic, r, "approve")
        assert "HOD" not in r.approval_route
        ln = r.lines.get()
        threshold = world.cfg.hod_approval_threshold
        client_for(world.stores).post(
            f"{API}/requests/{r.pk}/stores-edit/", {"lines": [{"id": ln.pk, "estimated_unit_price": str(threshold + 1)}]}, format="json"
        )
        r.refresh_from_db()
        assert r.approval_route[-1] == "HOD" and r.hod_required


# ---------------------------------------------------------------------------
# Equipment-linked inventory and maintenance
# ---------------------------------------------------------------------------
class TestLinksAndMaintenance:
    def test_links_and_suggested_lines(self, world):
        oil = make_item(world)
        o_ring = make_item(world, name="O-ring kit")
        stock_in(world, oil, "2", price="900.00")
        client_for(world.stores).post(
            f"{API}/stock/levels/", {"department_id": world.dept.pk, "item_id": oil.pk, "reorder_level": "3"}, format="json"
        )
        cl = client_for(world.oic)
        res = cl.post(f"{API}/item-links/", {"equipment_id": world.equipment.pk, "item_id": oil.pk, "usage": "CONSUMABLE", "typical_quantity": "2"}, format="json")
        assert res.status_code == 201, res.json()
        assert cl.post(f"{API}/item-links/", {"equipment_id": world.equipment.pk, "item_id": oil.pk}, format="json").json()["code"] == "duplicate_link"
        cl.post(f"{API}/item-links/", {"equipment_id": world.equipment.pk, "item_id": o_ring.pk, "usage": "SPARE"}, format="json")
        assert client_for(world.oic2).post(f"{API}/item-links/", {"equipment_id": world.equipment.pk, "item_id": oil.pk}, format="json").status_code == 403
        sugg = client_for(world.operator).get(f"{API}/equipment/{world.equipment.pk}/suggested-lines/").json()["results"]
        assert [x["item_id"] for x in sugg] == [oil.pk, o_ring.pk]
        assert sugg[0]["reorder_due"] is True and sugg[0]["central_stock"] == "2.000" and sugg[0]["last_unit_price"] == "900.00"
        assert sugg[1]["usage"] == "SPARE" and sugg[1]["central_stock"] == "0"
        assert client_for(world.other_oic).get(f"{API}/equipment/{world.equipment.pk}/suggested-lines/").status_code == 404

    def test_maintenance_with_parts_and_follow_up_request(self, world):
        oil = make_item(world)
        stock_in(world, oil, "5", price="400.00")
        start = timezone.now() - timedelta(hours=30)
        body = {"equipment_id": world.equipment.pk, "kind": "BREAKDOWN", "downtime_start": start.isoformat(),
                "downtime_end": (start + timedelta(hours=24)).isoformat(), "cause": "Pump seized", "action_taken": "Oil changed",
                "service_cost": "3000", "parts": [{"item_id": oil.pk, "quantity": "2"}]}
        assert client_for(world.oic).post(f"{API}/maintenance/", body, format="json").status_code == 403  # central store needs Stores
        res = client_for(world.stores).post(f"{API}/maintenance/", body, format="json")
        assert res.status_code == 201, res.json()
        rec = res.json()
        assert rec["downtime_hours"] == 24.0 and rec["parts_cost"] == "800.00" and rec["total_cost"] == "3800.00"
        assert central(world, oil) == Decimal("3.000") and rec["parts_used"][0]["reason_code"] == "CONSUMED_IN_REPAIR"
        no_parts = {k: v for k, v in body.items() if k != "parts"}
        assert client_for(world.oic).post(f"{API}/maintenance/", no_parts, format="json").status_code == 201
        assert client_for(world.oic2).post(f"{API}/maintenance/", no_parts, format="json").status_code == 403

        res = client_for(world.oic).post(
            f"{API}/maintenance/{rec['id']}/raise-request/", {"request_type": "REPAIR_MAINTENANCE", "lines": [line(price="5000")]}, format="json"
        )
        assert res.status_code == 201, res.json()
        pr = PurchaseRequest.objects.get(pk=res.json()["request"]["id"])
        assert pr.maintenance_record_id == rec["id"] and res.json()["submitted"] is True
        ov = client_for(world.oic).get(f"{API}/equipment/{world.equipment.pk}/overview/").json()
        assert ov["maintenance_totals"]["count"] == 2 and ov["can_record_maintenance"] is True
        assert pr.pk in [x["id"] for x in ov["open_requests"]]

    def test_back_to_functional_hook_links_request_and_maintenance(self, world):
        from iic_booking.equipment.disruption_procurement import procurement_options, raise_procurement_request
        from iic_booking.equipment.models import DisruptionEvent

        oil = make_item(world)
        start = timezone.now() - timedelta(days=2)
        event = DisruptionEvent.objects.create(
            equipment=world.equipment, disruption_type="UNDER_MAINTENANCE", scope="EQUIPMENT", start_at=start,
            end_at=timezone.now(), reason="Vacuum fault", action_taken="Pump serviced",
        )
        opts = procurement_options(world.oic, world.equipment)
        assert {"CONSUMABLE", "REPAIR_MAINTENANCE"} <= {x["value"] for x in opts["categories"]}
        assert opts["can_record_maintenance"] is True
        out = raise_procurement_request(event, world.oic, {
            "category": "CONSUMABLE",
            "items": [{"item_id": oil.pk, "name": "", "quantity": 3, "estimated_cost": 450}],
            "maintenance": {"kind": "BREAKDOWN", "service_cost": "2500"},
        })
        pr = PurchaseRequest.objects.get(pk=out["id"])
        rec = MaintenanceRecord.objects.get(disruption_event=event)
        assert pr.disruption_event_id == event.pk and pr.maintenance_record_id == rec.pk and out["maintenance_record"] == rec.number
        assert rec.cause == "Vacuum fault" and rec.downtime_start == start and rec.service_cost == Decimal("2500.00")
        assert pr.lines.get().item_id == oil.pk and pr.lines.get().description == oil.name
        event.refresh_from_db()
        assert event.procurement_request_ids == [pr.pk]


# ---------------------------------------------------------------------------
# Purchase mode, bills to Accounts, dashboard
# ---------------------------------------------------------------------------
class TestPurchaseModeAndAccounts:
    def _record(self, world, amount):
        return ProcurementRecord.objects.create(
            number=f"PRC-{uuid.uuid4().hex[:6]}", department=world.dept, financial_year="2026-27", title="Pump",
            approved_amount=Decimal(amount), estimated_amount=Decimal(amount), created_by=world.office,
        )

    def test_suggestion_and_recording(self, world):
        config_service.assign_role(world.admin, world.dept, world.office, R.OFFICE, [c.OfficePermission.PROCUREMENT, c.OfficePermission.INVOICES])
        cl = client_for(world.office)
        for amount, mode in (("40000", "DIRECT"), ("600000", "PURCHASE_COMMITTEE"), ("2000000", "LIMITED_TENDER"), ("9000000", "OPEN_TENDER")):
            rec = self._record(world, amount)
            assert cl.get(f"{API}/records/{rec.pk}/purchase-mode/").json()["suggested"] == mode
        url = f"{API}/records/{rec.pk}/purchase-mode/"
        assert cl.post(url, {"purchase_mode": "GEM"}, format="json").status_code == 400
        assert cl.post(url, {"purchase_mode": "DIRECT"}, format="json").status_code == 400  # deviation needs a reason
        res = cl.post(url, {"purchase_mode": "GEM", "gem_reference": "GEMC-511687712345"}, format="json")
        assert res.status_code == 200 and res.json()["purchase_mode"] == "GEM"

    def test_forward_bill_to_accounts(self, world):
        accounts = make_user(user_type=UserType.FINANCE, name="Accounts", department=world.dept)
        config_service.assign_role(world.admin, world.dept, accounts, R.ACCOUNTS)
        rec = self._record(world, "40000")
        inv = Invoice.objects.create(
            department=world.dept, procurement_record=rec, invoice_number="INV-1", invoice_date=TODAY(),
            total_amount=Decimal("40000"), recorded_by=world.office,
        )
        assert client_for(world.operator).post(f"{API}/invoices/{inv.pk}/forward/", {}, format="json").status_code in (403, 404)
        res = client_for(world.stores).post(f"{API}/invoices/{inv.pk}/forward/", {"note": "GRN done"}, format="json")
        assert res.status_code == 200 and res.json()["forwarded_to_accounts_at"]
        assert client_for(world.stores).post(f"{API}/invoices/{inv.pk}/forward/", {}, format="json").json()["code"] == "already_forwarded"
        bills = client_for(accounts).get(f"{API}/accounts/bills/").json()
        assert [b["id"] for b in bills["results"]] == [inv.pk]
        dash = client_for(accounts).get(f"{API}/dashboard/", {"department_id": world.dept.pk}).json()
        assert dash["accounts"]["bills_pending"] == 1

    def test_dashboard_ageing_and_menus(self, world):
        r = new_request(world.operator, equipment=world.equipment, submit=True)
        PurchaseRequest.objects.filter(pk=r.pk).update(stage_entered_at=timezone.now() - timedelta(days=9))
        dash = client_for(world.oic).get(f"{API}/dashboard/", {"department_id": world.dept.pk}).json()
        assert dash["my_queue_ageing"]["buckets"]["8-15"] == 1 and dash["pending_for_me"][0]["stage_age_days"] == 9
        boot = client_for(world.stores).get(f"{API}/bootstrap/").json()
        assert boot["menus"]["registers"] and boot["menus"]["register_import"] and boot["menus"]["verification"]
        boot = client_for(world.operator).get(f"{API}/bootstrap/").json()
        assert boot["menus"]["registers"] and not boot["menus"]["register_import"] and not boot["menus"]["accounts"]
