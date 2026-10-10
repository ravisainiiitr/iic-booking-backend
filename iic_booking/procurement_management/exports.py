"""CSV / XLSX / PDF exports (tabular reports and the proposal document)."""

from __future__ import annotations

import csv
import io
import re

from django.http import HttpResponse
from django.utils import timezone

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _cell(value) -> str:
    """Neutralise spreadsheet formula injection for text cells."""
    text = "" if value is None else str(value)
    if text and text[0] in FORMULA_PREFIXES and not re.fullmatch(r"-?\d+(\.\d+)?", text):
        return "'" + text
    return text


def _filename(base: str, ext: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", base).strip("_") or "export"
    return f"{safe}_{timezone.localdate():%Y%m%d}.{ext}"


SERIAL_HEADERS = ("#", "S.No", "S.No.")
LONG_TEXT = 60


def _with_serial(headers: list[str], rows: list[list]) -> tuple[list[str], list[list]]:
    if headers and str(headers[0]).strip() in SERIAL_HEADERS:
        return ["S.No.", *headers[1:]], rows
    return ["S.No.", *headers], [[i, *row] for i, row in enumerate(rows, 1)]


def _xlsx(title: str, subtitle: str, headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import Border
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill
    from openpyxl.styles import Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]+", "-", title)[:31] or "Report"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    if subtitle:
        ws.append([subtitle])
    ws.append([f"Generated {timezone.localtime():%d %b %Y %H:%M}"])
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="153F79", end_color="153F79", fill_type="solid")
        cell.alignment = centre
        cell.border = border
    widths = [len(str(h)) for h in headers]
    for row in rows:
        ws.append([v if isinstance(v, (int, float)) else _cell(v) for v in row])
        for index, cell in enumerate(ws[ws.max_row]):
            text = "" if cell.value is None else str(cell.value)
            cell.alignment = left if len(text) > LONG_TEXT else centre
            cell.border = border
            if index < len(widths):
                widths[index] = max(widths[index], min(len(text), 50))
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = max(8, width + 3)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_response(fmt: str, title: str, headers: list[str], rows: list[list], *, subtitle: str = "") -> HttpResponse:
    fmt = (fmt or "csv").lower()
    headers, rows = _with_serial(list(headers), list(rows))
    if fmt == "xlsx":
        resp = HttpResponse(_xlsx(title, subtitle, headers, rows),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "xlsx")}"'
        return resp
    if fmt == "pdf":
        resp = HttpResponse(_pdf(title, subtitle, headers, rows), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "pdf")}"'
        return resp
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([title])
    if subtitle:
        writer.writerow([subtitle])
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_cell(v) for v in row])
    resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "csv")}"'
    return resp


def _pdf(title: str, subtitle: str, headers: list[str], rows: list[list], *, extra_story=None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24, title=title)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    small_center = small.clone("small_center", alignment=1)
    head = small.clone("head", alignment=1, textColor=colors.white)
    story = [Paragraph("Indian Institute of Technology Roorkee", styles["Title"]), Paragraph(title, styles["Heading2"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["BodyText"]))
    story.append(Paragraph(f"Generated {timezone.localtime():%d %b %Y %H:%M}", small))
    story.append(Spacer(1, 8))
    def esc(v) -> str:
        return str("" if v is None else v).replace("&", "&amp;").replace("<", "&lt;").replace("₹", "Rs. ")

    data = [[Paragraph(f"<b>{esc(h)}</b>", head) for h in headers]]
    for row in rows:
        data.append([Paragraph(esc(v), small if len(esc(v)) > LONG_TEXT else small_center) for v in row])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#153f79")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")]),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(table)
    for item in extra_story or []:
        story.append(item)
    doc.build(story)
    return buf.getvalue()


def proposal_pdf_response(p) -> HttpResponse:
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, Spacer

    reqs = p.requirements.exclude(status__in=["REMOVED", "MERGED"]).select_related("equipment", "laboratory").order_by("number")
    headers = ["S.No.", "Number", "Description", "Lab / Equipment", "Qty", "UoM", "Unit cost (₹)", "Estimate (₹)", "Approved (₹)", "Priority"]
    rows = []
    for i, r in enumerate(reqs, 1):
        where = (r.equipment.name if r.equipment_id else "") or (r.laboratory.name if r.laboratory_id else "")
        rows.append([i, r.number, r.description, where, r.quantity, r.uom, r.estimated_unit_cost, r.estimated_total,
                     r.approved_amount if r.approved_amount is not None else "", r.get_priority_display()])
    styles = getSampleStyleSheet()
    extra = [
        Spacer(1, 10),
        Paragraph(f"<b>Total estimate:</b> Rs. {p.total_amount}", styles["BodyText"]),
        Spacer(1, 30),
        Paragraph("Prepared by (Office) ____________________ &nbsp;&nbsp;&nbsp; Approved by (HOD) ____________________", styles["BodyText"]),
    ]
    subtitle = f"{p.department.name} — {p.get_funding_type_display()} {p.financial_year} — {p.number} — Status: {p.get_status_display()}"
    pdf = _pdf(f"Proposal: {p.title}", subtitle, headers, rows, extra_story=extra)
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(p.number, "pdf")}"'
    return resp
