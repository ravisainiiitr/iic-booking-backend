"""Bulk import of existing Major / Minor asset registers from Excel or CSV.

Flow: download the template → fill one row per register entry → upload for a *preview* (every row validated,
duplicates on (register, page, serial) flagged against the database and within the file) → *commit* (the file is
sent again; rows are re-validated and created in one transaction, optionally skipping rows with errors).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone

from . import access, audit
from . import constants as c
from .errors import ProcurementError
from .fy import fy_label, is_valid_fy_label
from .models import Asset, AssetRegister, ItemCategory
from .numbering import next_number

P = c.OfficePermission
MAX_ROWS = 5000
MAX_BYTES = 10 * 1024 * 1024

# (header, key, required, note)
COLUMNS = [
    ("Register Code", "register_code", True, "Code on the register cover, e.g. MAJ-1. Created on import if missing and allowed."),
    ("Register Type", "register_type", True, "MAJOR, MINOR or LIMITED_LIFE."),
    ("Page No", "page", True, "Page number in the register (1, 2, …)."),
    ("Serial No", "serial", True, "Serial number of the entry on that page."),
    ("Entry Date", "entry_date", False, "Date written against the entry (YYYY-MM-DD or DD-MM-YYYY)."),
    ("Description", "description", True, "Item description as written in the register."),
    ("Category Code", "category_code", False, "Item category code; defaults to the category for the register type."),
    ("Quantity", "quantity", False, "Number of units covered by this entry (default 1)."),
    ("Make", "make", False, ""),
    ("Model", "model_number", False, ""),
    ("Manufacturer Serial No", "serial_number", False, "Serial number printed on the item."),
    ("Asset Tag", "asset_tag", False, "Existing tag / sticker number. Leave blank to auto-generate."),
    ("Equipment Code", "equipment_code", False, "Booking-system equipment code this asset belongs to."),
    ("Parent Asset Tag", "parent_tag", False, "For accessories: the asset tag of the main asset (in the system or in this file)."),
    ("Laboratory Code", "laboratory_code", False, ""),
    ("Location", "location", False, "Room / location."),
    ("Custodian Email", "custodian_email", False, "Institute e-mail of the custodian."),
    ("Supplier Name", "supplier_name", False, ""),
    ("Supplier GSTIN", "supplier_gstin", False, "Matched to a vendor in the vendor master when present."),
    ("PO Number", "po_number", False, ""),
    ("PO Date", "po_date", False, ""),
    ("Invoice Number", "invoice_number", False, ""),
    ("Invoice Date", "invoice_date", False, ""),
    ("Cost", "cost", False, "Cost per unit in rupees (numbers only)."),
    ("Funding Source", "funding_source", False, "e.g. Plan grant, project, institute."),
    ("Project Code", "project_code", False, ""),
    ("Financial Year", "financial_year", False, "e.g. 2019-20. Derived from the purchase / entry date when blank."),
    ("Warranty Until", "warranty_until", False, ""),
    ("AMC Until", "amc_until", False, ""),
    ("Installation Date", "installation_date", False, ""),
    ("Condition", "condition", False, "NEW, GOOD, FAIR, POOR or UNSERVICEABLE."),
    ("Status", "status", False, "ACTIVE (default), IN_USE, IN_STORE, UNDER_REPAIR, DAMAGED, LOST, CONDEMNED, DISPOSED, RETIRED."),
    ("Remarks", "remarks", False, ""),
    ("Legacy Ref", "legacy_ref", False, "Old stock / dead-stock number, if any."),
]
HEADER_KEY = {re.sub(r"[^a-z0-9]", "", h.lower()): k for h, k, _, _ in COLUMNS}
REQUIRED = [k for _, k, req, _ in COLUMNS if req]
IMPORT_STATUSES = [s for s in c.AssetStatus.values if s not in c.ASSET_TRANSFER_STATUSES]
ASSET_REGISTER_TYPES = [c.RegisterType.MAJOR, c.RegisterType.MINOR, c.RegisterType.LIMITED_LIFE]
EXAMPLE = {
    "register_code": "MAJ-1", "register_type": "MAJOR", "page": "12", "serial": "3", "entry_date": "2019-08-14",
    "description": "Field emission scanning electron microscope", "quantity": "1", "make": "Carl Zeiss",
    "model_number": "Gemini 500", "serial_number": "ZS-50012", "equipment_code": "FESEM-01",
    "location": "Room 112", "supplier_name": "Carl Zeiss India", "po_number": "IITR/PUR/2019/123",
    "po_date": "2019-06-01", "invoice_number": "INV-889", "invoice_date": "2019-07-20", "cost": "25000000",
    "funding_source": "Plan grant", "condition": "GOOD", "status": "ACTIVE",
}


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------
def template_response(fmt: str = "xlsx") -> HttpResponse:
    headers = [h for h, _, _, _ in COLUMNS]
    example = [EXAMPLE.get(k, "") for _, k, _, _ in COLUMNS]
    if (fmt or "xlsx").lower() == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(headers)
        w.writerow(example)
        resp = HttpResponse("\ufeff" + buf.getvalue(), content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = 'attachment; filename="asset_register_import_template.csv"'
        return resp
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(headers)
    ws.append(example)
    req_fill = PatternFill("solid", fgColor="FDE2E1")
    for idx, (_, _, req, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if req:
            cell.fill = req_fill
        ws.column_dimensions[cell.column_letter].width = max(14, min(40, len(cell.value) + 4))
    for cell in ws[2]:
        cell.font = Font(italic=True, color="808080")
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Instructions")
    notes.append(["Column", "Required", "Notes"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for h, _, req, note in COLUMNS:
        notes.append([h, "Yes" if req else "", note])
    notes.append([])
    notes.append(["Row 2 of the Assets sheet is an example — delete it before uploading."])
    notes.append(["(Register Code, Page No, Serial No) must be unique; duplicates are flagged in the preview."])
    notes.append([f"At most {MAX_ROWS} rows per file. Dates: YYYY-MM-DD or DD-MM-YYYY."])
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["C"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="asset_register_import_template.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def read_rows(upload) -> list[dict]:
    if upload is None:
        raise ProcurementError("Attach the filled template.", code="file_required", field="file")
    if upload.size > MAX_BYTES:
        raise ProcurementError("The file is larger than 10 MB.", code="file_too_large", field="file")
    name = (upload.name or "").lower()
    raw = upload.read()
    if name.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        table = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ProcurementError("The file could not be read as an Excel workbook.", code="invalid_file", field="file")
        ws = wb["Assets"] if "Assets" in wb.sheetnames else wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ProcurementError("Upload an .xlsx or .csv file.", code="invalid_file", field="file")
    while table and not any(_text(v) for v in table[0]):
        table.pop(0)
    if not table:
        raise ProcurementError("The file is empty.", code="empty_file", field="file")
    keys = [HEADER_KEY.get(_norm_header(h)) for h in table[0]]
    missing = [h for h, k, req, _ in COLUMNS if req and k not in keys]
    if missing:
        raise ProcurementError(f"Missing columns: {', '.join(missing)}. Use the template.", code="invalid_template", missing=missing)
    rows = []
    for idx, values in enumerate(table[1:], start=2):
        row = {k: _text(v) for k, v in zip(keys, values) if k}
        if not any(row.values()):
            continue
        row["_row"] = idx
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise ProcurementError(f"At most {MAX_ROWS} rows per file.", code="too_many_rows")
    return rows


def _date(v: str):
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(v[:10], fmt).date()
        except ValueError:
            continue
    raise ValueError(v)


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
@dataclass
class RowResult:
    row: int
    data: dict
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    duplicate_of: dict | None = None
    values: dict = field(default_factory=dict)

    @property
    def status(self) -> str:
        if self.duplicate_of:
            return "DUPLICATE"
        if self.errors:
            return "ERROR"
        return "WARNING" if self.warnings else "OK"

    def as_dict(self) -> dict:
        return {
            "row": self.row,
            "status": self.status,
            "errors": self.errors,
            "warnings": self.warnings,
            "duplicate_of": self.duplicate_of,
            "register_ref": f"{self.data.get('register_code', '')} / p.{self.data.get('page', '')} / s.{self.data.get('serial', '')}",
            "description": self.data.get("description", ""),
            "asset_tag": self.data.get("asset_tag", ""),
            "equipment_code": self.data.get("equipment_code", ""),
            "cost": self.data.get("cost", ""),
        }


class _Lookups:
    def __init__(self, dept):
        from iic_booking.equipment.models import Equipment
        from iic_booking.sync.models import Laboratory
        from iic_booking.users.models import User

        from .models import Vendor

        self.dept = dept
        self.registers = {r.code.upper(): r for r in AssetRegister.objects.filter(department=dept, is_archived=False)}
        cats = list(ItemCategory.objects.filter(department=dept, active=True, is_asset=True))
        self.categories = {x.code.upper(): x for x in cats}
        self.default_category = {}
        for rtype, nature in c.REGISTER_CATEGORY_NATURE.items():
            match = [x for x in cats if x.nature == nature]
            if match:
                self.default_category[rtype] = match[0]
        self.equipment = {
            (e.code or "").upper(): e for e in Equipment.objects.filter(internal_department=dept).only("equipment_id", "code", "name")
            if e.code
        }
        self.labs = {(lab.code or "").upper(): lab for lab in Laboratory.objects.filter(department=dept, is_active=True) if lab.code}
        self._users = User
        self._user_cache: dict[str, object] = {}
        self.vendors = {v.gstin.upper(): v for v in Vendor.objects.filter(department=dept, is_archived=False).exclude(gstin="")}
        self.existing_tags = {
            t.upper(): pk for pk, t in Asset.objects.filter(department=dept, is_archived=False).exclude(asset_tag="").values_list("id", "asset_tag")
        }

    def user(self, email: str):
        key = email.lower()
        if key not in self._user_cache:
            self._user_cache[key] = self._users.objects.filter(email__iexact=email, is_active=True).first()
        return self._user_cache[key]


def _validate(rows: list[dict], lk: _Lookups, *, create_registers: bool) -> list[RowResult]:
    results: list[RowResult] = []
    seen_entries: dict[tuple, int] = {}
    seen_tags: dict[str, int] = {}
    file_tags = {r.get("asset_tag", "").upper() for r in rows if r.get("asset_tag")}
    for raw in rows:
        res = RowResult(row=raw["_row"], data={k: v for k, v in raw.items() if k != "_row"})
        d, v = res.data, res.values
        for key in REQUIRED:
            if not d.get(key):
                header = next(h for h, k, _, _ in COLUMNS if k == key)
                res.errors.append(f"{header} is required.")
        rtype = d.get("register_type", "").upper().replace(" ", "_").replace("-", "_")
        if rtype in ("LIMITEDLIFE",):
            rtype = c.RegisterType.LIMITED_LIFE
        if rtype and rtype not in ASSET_REGISTER_TYPES:
            res.errors.append("Register Type must be MAJOR, MINOR or LIMITED_LIFE.")
        code = d.get("register_code", "").upper().replace(" ", "-")
        reg = lk.registers.get(code)
        if code and reg is None and not create_registers:
            res.errors.append(f"Register {code} does not exist (tick “create missing registers” or add it first).")
        if reg is not None and rtype and reg.register_type != rtype:
            res.errors.append(f"Register {code} is a {reg.get_register_type_display()} register.")
        v["register"], v["register_code"], v["register_type"] = reg, code, rtype
        try:
            page = int(Decimal(d.get("page") or "0"))
            if page < 1:
                raise InvalidOperation
            if reg is not None and reg.total_pages and page > reg.total_pages:
                res.errors.append(f"{code} has only {reg.total_pages} pages.")
            v["page"] = page
        except (InvalidOperation, ValueError):
            if d.get("page"):
                res.errors.append("Page No must be a whole number.")
        serial = d.get("serial", "")[:20]
        v["serial"] = serial
        if code and v.get("page") and serial:
            key = (code, v["page"], serial.upper())
            if key in seen_entries:
                res.duplicate_of = {"row": seen_entries[key]}
                res.errors.append(f"Same register / page / serial as row {seen_entries[key]}.")
            else:
                seen_entries[key] = res.row
                if reg is not None:
                    hit = Asset.objects.filter(
                        register=reg, register_page=v["page"], register_serial__iexact=serial, is_archived=False
                    ).values("id", "number", "description").first()
                    if hit:
                        res.duplicate_of = {"asset_id": hit["id"], "asset_number": hit["number"], "description": hit["description"]}
                        res.errors.append(f"Already registered as {hit['number']}.")
        cat = lk.categories.get(d.get("category_code", "").upper()) if d.get("category_code") else lk.default_category.get(rtype)
        if d.get("category_code") and cat is None:
            res.errors.append(f"Unknown asset category {d['category_code']}.")
        elif cat is None and rtype:
            res.errors.append("No asset category for this register type — give a Category Code.")
        v["category"] = cat
        try:
            qty = int(Decimal(d.get("quantity") or "1"))
            if qty < 1:
                raise InvalidOperation
            v["quantity"] = qty
        except (InvalidOperation, ValueError):
            res.errors.append("Quantity must be a whole number of at least 1.")
        cost_raw = (d.get("cost") or "").replace(",", "").replace("₹", "").strip()
        try:
            v["cost"] = Decimal(cost_raw or "0").quantize(Decimal("0.01"))
            if v["cost"] < 0:
                raise InvalidOperation
            if not cost_raw:
                res.warnings.append("No cost given — recorded as ₹0.00.")
        except (InvalidOperation, ValueError):
            res.errors.append("Cost must be a number.")
        for key in ("entry_date", "po_date", "invoice_date", "warranty_until", "amc_until", "installation_date"):
            try:
                v[key] = _date(d.get(key, ""))
            except ValueError:
                header = next(h for h, k, _, _ in COLUMNS if k == key)
                res.errors.append(f"{header} is not a valid date.")
        fy = d.get("financial_year", "")
        if fy and not is_valid_fy_label(fy):
            res.errors.append("Financial Year must look like 2019-20.")
        base_day = v.get("invoice_date") or v.get("po_date") or v.get("entry_date")
        v["financial_year"] = fy or (fy_label(base_day) if base_day else "")
        cond = d.get("condition", "").upper()
        if cond and cond not in c.AssetCondition.values:
            res.errors.append("Condition must be NEW, GOOD, FAIR, POOR or UNSERVICEABLE.")
        v["condition"] = cond
        status = (d.get("status") or "ACTIVE").upper().replace(" ", "_")
        if status not in IMPORT_STATUSES:
            res.errors.append(f"Unknown status {d.get('status')}.")
        v["status"] = status
        eq = None
        if d.get("equipment_code"):
            eq = lk.equipment.get(d["equipment_code"].upper())
            if eq is None:
                res.errors.append(f"No equipment with code {d['equipment_code']} in this department.")
        v["equipment"] = eq
        lab = None
        if d.get("laboratory_code"):
            lab = lk.labs.get(d["laboratory_code"].upper())
            if lab is None:
                res.errors.append(f"No laboratory with code {d['laboratory_code']} in this department.")
        v["laboratory"] = lab
        cust = None
        if d.get("custodian_email"):
            cust = lk.user(d["custodian_email"])
            if cust is None:
                res.warnings.append(f"Custodian {d['custodian_email']} not found — left blank.")
        v["custodian"] = cust
        v["vendor"] = lk.vendors.get(d.get("supplier_gstin", "").upper()) if d.get("supplier_gstin") else None
        tag = d.get("asset_tag", "")[:120]
        if tag:
            if tag.upper() in lk.existing_tags:
                res.errors.append(f"Asset tag {tag} is already in use.")
            elif tag.upper() in seen_tags:
                res.errors.append(f"Asset tag {tag} repeats row {seen_tags[tag.upper()]}.")
            else:
                seen_tags[tag.upper()] = res.row
        v["asset_tag"] = tag
        parent = d.get("parent_tag", "")
        if parent and parent.upper() not in lk.existing_tags and parent.upper() not in file_tags:
            res.errors.append(f"Parent asset tag {parent} not found in the system or this file.")
        if parent and tag and parent.upper() == tag.upper():
            res.errors.append("An asset cannot be its own parent.")
        v["parent_tag"] = parent
        if d.get("serial_number") and Asset.objects.filter(
            department=lk.dept, serial_number__iexact=d["serial_number"], is_archived=False
        ).exists():
            res.warnings.append(f"Manufacturer serial {d['serial_number']} already exists on another asset.")
        results.append(res)
    return results


def _department(scope, data):
    from .registers import _cfg

    dept = access.pick_department(scope, data.get("department_id"))
    _cfg(dept.pk)
    scope.require_perm(dept.pk, P.ASSETS)
    return dept


def preview(scope, data: dict, upload) -> dict:
    dept = _department(scope, data)
    from .api import parse_bool

    rows = read_rows(upload)
    results = _validate(rows, _Lookups(dept), create_registers=parse_bool(data.get("create_registers") or ""))
    counts = {"OK": 0, "WARNING": 0, "ERROR": 0, "DUPLICATE": 0}
    for r in results:
        counts[r.status] += 1
    new_regs = sorted({r.values["register_code"] for r in results if r.values.get("register_code") and r.values.get("register") is None})
    return {
        "department_id": dept.pk,
        "total": len(results),
        "counts": counts,
        "importable": counts["OK"] + counts["WARNING"],
        "new_registers": new_regs,
        "rows": [r.as_dict() for r in results],
    }


@transaction.atomic
def commit(scope, data: dict, upload, *, request=None) -> dict:
    from .api import parse_bool
    from .assets import _history
    from .registers import ensure_tag, save_entry

    dept = _department(scope, data)
    create_registers = parse_bool(data.get("create_registers") or "")
    skip_errors = parse_bool(data.get("skip_errors") or "")
    rows = read_rows(upload)
    lk = _Lookups(dept)
    results = _validate(rows, lk, create_registers=create_registers)
    bad = [r for r in results if r.errors]
    if bad and not skip_errors:
        raise ProcurementError(
            f"{len(bad)} row(s) have errors. Fix them or choose to skip them.", code="import_errors",
            rows=[r.as_dict() for r in bad[:200]],
        )
    good = [r for r in results if not r.errors]
    new_codes: list[str] = []
    for r in good:
        code = r.values["register_code"]
        if r.values["register"] is None:
            reg = lk.registers.get(code)
            if reg is None:
                reg = AssetRegister.objects.create(
                    department=dept, register_type=r.values["register_type"], code=code,
                    name=f"{c.RegisterType(r.values['register_type']).label} register {code}", created_by=scope.user,
                )
                lk.registers[code] = reg
                new_codes.append(code)
                audit.record(scope.user, "register.created", reg, new={"code": code, "source": "import"}, request=request)
            r.values["register"] = reg
    # Main assets first so accessories in the same file can point at them.
    good.sort(key=lambda r: 1 if r.values["parent_tag"] else 0)
    tag_to_asset: dict[str, Asset] = {}
    created = 0
    for r in good:
        v, d = r.values, r.data
        parent = None
        if v["parent_tag"]:
            parent = tag_to_asset.get(v["parent_tag"].upper())
            if parent is None:
                parent = Asset.objects.filter(pk=lk.existing_tags.get(v["parent_tag"].upper())).first()
        asset = Asset(
            number=next_number(c.NumberPrefix.ASSET),
            department=dept,
            laboratory=v["laboratory"],
            equipment=v["equipment"] or (parent.equipment if parent else None),
            category=v["category"],
            description=d["description"][:255],
            make=d.get("make", "")[:120],
            model_number=d.get("model_number", "")[:120],
            serial_number=d.get("serial_number", "")[:120],
            asset_tag=v["asset_tag"],
            vendor=v["vendor"],
            purchase_date=v.get("invoice_date") or v.get("po_date"),
            cost=v["cost"],
            is_capitalized=v["register_type"] == c.RegisterType.MAJOR,
            financial_year=v["financial_year"],
            warranty_until=v.get("warranty_until"),
            location=d.get("location", "")[:255],
            custodian=v["custodian"],
            status=v["status"],
            remarks=d.get("remarks", "")[:5000],
            created_by=scope.user,
            register=v["register"],
            register_page=v["page"],
            register_serial=v["serial"],
            register_entry_date=v.get("entry_date"),
            legacy_ref=d.get("legacy_ref", "")[:120],
            parent=parent,
            quantity=v["quantity"],
            supplier_name=(d.get("supplier_name") or (v["vendor"].name if v["vendor"] else ""))[:255],
            po_number=d.get("po_number", "")[:80],
            po_date=v.get("po_date"),
            invoice_number=d.get("invoice_number", "")[:80],
            invoice_date=v.get("invoice_date"),
            funding_source=d.get("funding_source", "")[:255],
            project_code=d.get("project_code", "")[:80],
            installation_date=v.get("installation_date"),
            amc_until=v.get("amc_until"),
            condition=v["condition"],
        )
        save_entry(asset)
        ensure_tag(asset)
        tag_to_asset[asset.asset_tag.upper()] = asset
        _history(asset, "", asset.status, f"Imported from register {asset.register_ref} (row {r.row})", scope.user)
        created += 1
    audit.record(
        scope.user, "asset.imported", dept, department=dept,
        new={"created": created, "skipped": len(bad), "file": getattr(upload, "name", "")[:120]}, request=request,
    )
    return {
        "created": created,
        "skipped": len(bad),
        "skipped_rows": [r.as_dict() for r in bad[:200]],
        "registers_created": new_codes,
        "imported_at": timezone.now().isoformat(),
    }
