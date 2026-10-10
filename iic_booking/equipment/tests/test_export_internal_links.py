"""Exports that contain both a summary and the details jump within the file instead of opening the portal."""

import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from iic_booking.equipment import report_exports

from .test_booking_export_full_inputs import add_rich_inputs
from .test_booking_list_export import _export
from .test_booking_list_export import world  # noqa: F401 - fixture


def _link_annotations(page):
    for ref in page.get("/Annots") or []:
        annot = ref.get_object()
        if annot.get("/Subtype") == "/Link":
            yield annot


def _outline(reader, items=None):
    """[(title, page index)] of every outline entry, depth first."""
    out = []
    for item in reader.outline if items is None else items:
        if isinstance(item, list):
            out.extend(_outline(reader, item))
        else:
            out.append((item.title, reader.get_destination_page_number(item)))
    return out


def _target_page(reader, annot) -> int:
    page_ref = annot["/Dest"].get_object()[0]
    return next(i for i, page in enumerate(reader.pages) if page.indirect_reference.idnum == page_ref.idnum)


@pytest.mark.django_db
def test_booking_pdf_summary_ids_jump_to_their_details_card(world, settings):
    settings.FRONTEND_URL = "https://equip.example.org"
    add_rich_inputs(world)
    res = _export(world.f.client_for(world.admin), "pdf", view="staff")
    assert res.status_code == 200
    reader = PdfReader(io.BytesIO(res.content))
    booking_ids = {b.virtual_booking_id for b in (world.b1, world.b2, world.b3)}

    outline = _outline(reader)
    assert [title for title, _ in outline[:2]] == ["Summary", "Booking details"]
    assert outline[0][1] == 0
    cards = [(title, page) for title, page in outline[2:]]
    assert len(cards) == len(booking_ids)
    for title, page in cards:
        booking_id = title.split(". ", 1)[1]
        assert booking_id in booking_ids
        text = reader.pages[page].extract_text()
        assert booking_id in text and "Back to summary" in text
    # Destinations are registered on the page being emitted, not stuck on the first page.
    assert min(page for _, page in cards) >= 1

    summary_links = list(_link_annotations(reader.pages[0]))
    assert len(summary_links) == len(booking_ids)
    assert all("/Dest" in a and "/A" not in a for a in summary_links)
    assert sorted(_target_page(reader, a) for a in summary_links) == sorted(page for _, page in cards)

    card_links = [a for page in {p for _, p in cards} for a in _link_annotations(reader.pages[page])]
    assert all(_target_page(reader, a) == 0 for a in card_links if "/Dest" in a)
    uris = [a["/A"]["/URI"] for a in card_links if "/A" in a]
    assert uris and all(u.startswith("https://equip.example.org/booking-management?expand=") for u in uris)


@pytest.mark.django_db
def test_booking_xlsx_detail_sheets_link_to_the_bookings_row(world, settings):
    settings.FRONTEND_URL = "https://equip.example.org"
    b1 = add_rich_inputs(world)
    wb = load_workbook(io.BytesIO(_export(world.f.client_for(world.admin), "xlsx", view="staff").content))
    bookings = wb["Bookings"]
    header = [c.value for c in bookings[1]]
    id_col = header.index("Booking ID") + 1

    for sheet in ("Sample sets", "Charges"):
        cell = next(r[0] for r in wb[sheet].iter_rows(min_row=2) if r[0].value == b1.virtual_booking_id)
        assert cell.hyperlink.target is None
        target_sheet, target_cell = cell.hyperlink.location.split("!")
        assert target_sheet == "'Bookings'"
        assert bookings[target_cell].value == b1.virtual_booking_id
        assert bookings[target_cell].column == id_col

    tables = wb["Input tables"]
    linked = [c for row in tables.iter_rows() for c in row[:1] if c.hyperlink is not None]
    assert linked and all(c.hyperlink.location.startswith("'Bookings'!") for c in linked)

    own_row = next(r for r in bookings.iter_rows(min_row=2) if r[id_col - 1].value == b1.virtual_booking_id)
    assert own_row[id_col - 1].hyperlink.target.startswith("https://equip.example.org/booking-management?expand=")

    filters = {r[0].value: r[0].hyperlink for r in wb["Filters"].iter_rows() if r and r[0].value}
    assert filters["Charges"].location == "'Charges'!A1"


def _equipment_report_data():
    def equipment(code, name):
        return {
            "name": name, "code": code, "officers_in_charge": [], "lab_operators": [], "slot_window_display": "",
            "booked_hours": 5, "no_booking_hours": 3, "user_ratings": {},
        }

    return {
        "report_header": {"report_title": "Equipment Performance Report"},
        "summary": {},
        "equipment": [equipment("XRD1", "Alpha XRD"), equipment("TGA1", "Beta TGA & DSC")],
        "financial": {"revenue_by_equipment": [
            {"equipment__code": "TGA1", "equipment__name": "Beta TGA & DSC", "count": 2, "total": 500},
            {"equipment__code": "GONE", "equipment__name": "Not in this report", "count": 1, "total": 10},
        ]},
        "utilization_pie": [],
    }


def test_equipment_report_pdf_links_equipment_to_its_section(monkeypatch):
    monkeypatch.setattr(report_exports, "get_equipment_report_data", lambda **_: _equipment_report_data())
    reader = PdfReader(io.BytesIO(report_exports.build_report_pdf()))
    outline = dict(_outline(reader))
    assert "Per-equipment performance" in outline and "Revenue by equipment" in outline
    alpha, beta = outline["Alpha XRD (XRD1)"], outline["Beta TGA & DSC (TGA1)"]
    assert "Alpha XRD" in reader.pages[alpha].extract_text()
    assert "Beta TGA & DSC" in reader.pages[beta].extract_text()
    assert beta > alpha

    links = [a for page in reader.pages for a in _link_annotations(page)]
    # One revenue row (the other equipment is not in the report) and two utilization headings; no web links.
    assert len(links) == 3 and all("/Dest" in a for a in links)
    revenue_page = outline["Revenue by equipment"]
    revenue_links = [a for a in _link_annotations(reader.pages[revenue_page])]
    assert _target_page(reader, revenue_links[0]) == beta
    assert sorted(_target_page(reader, a) for a in links) == sorted([alpha, beta, beta])


def test_equipment_report_xlsx_links_between_sheets(monkeypatch):
    monkeypatch.setattr(report_exports, "get_equipment_report_data", lambda **_: _equipment_report_data())
    wb = load_workbook(io.BytesIO(report_exports.build_report_excel()))
    main = wb["Equipment Report"]
    revenue = wb["Revenue by Equipment"]
    utilization = wb["Per-Equipment Utilization"]

    linked, unlinked = revenue["A5"], revenue["A6"]
    sheet, cell = linked.hyperlink.location.split("!")
    assert sheet == "'Equipment Report'" and main[cell].value == "Beta TGA & DSC"
    assert unlinked.hyperlink is None

    main_cell = main[cell]
    sheet, block = main_cell.hyperlink.location.split("!")
    assert sheet == "'Per-Equipment Utilization'" and utilization[block].value == "Beta TGA & DSC (TGA1)"
    back = utilization[block].hyperlink.location
    assert back == f"'Equipment Report'!{cell}"
