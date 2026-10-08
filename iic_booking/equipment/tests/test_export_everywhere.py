"""GET /api/exports/<report>/: same rows and scoping as each list endpoint, in xlsx / csv / pdf."""

import csv
import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from iic_booking.equipment.models import BookingAttemptLog
from iic_booking.equipment.models import BookingAttemptOutcome
from iic_booking.equipment.models import EquipmentManager
from iic_booking.equipment.models import UrgentBookingRequest
from iic_booking.equipment.models import UrgentBookingRequestType
from iic_booking.equipment.models import WaitlistEntry
from iic_booking.users.models.user_type import UserType
from iic_booking.users.tests.factories import UserFactory


def _staff(user_type, **kw):
    return UserFactory(user_type=user_type, admin_approved=True, **kw)


def _export(client, report, export_format="csv", **params):
    return client.get(f"/api/exports/{report}/", {"export_format": export_format, **params})


def _csv(response):
    assert response.status_code == 200, getattr(response, "data", response.content[:300])
    assert response.content.startswith("\ufeff".encode())
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def _column(rows, header):
    index = rows[0].index(header)
    return [r[index] for r in rows[1:]]


@pytest.fixture
def world(egs_factory):
    f = egs_factory
    eq_a = f.equipment(name="Alpha XRD")
    eq_b = f.equipment(name="Beta TGA")
    oic_a = _staff(UserType.MANAGER, department=f.department)
    EquipmentManager.objects.create(equipment=eq_a, manager=oic_a)
    admin = _staff(UserType.ADMIN, is_staff=True)
    alice = f.student()
    alice.name = "Alice Rao"
    alice.save(update_fields=["name"])
    bob = f.student()
    bob.name = '=HYPERLINK("http://x")'
    bob.save(update_fields=["name"])
    logs = [
        BookingAttemptLog.objects.create(user=alice, equipment=eq_a, outcome=BookingAttemptOutcome.FAILED,
                                         failure_reason="No slots available"),
        BookingAttemptLog.objects.create(user=bob, equipment=eq_b, outcome=BookingAttemptOutcome.FAILED,
                                         failure_reason="Weekly limit reached"),
        BookingAttemptLog.objects.create(user=alice, equipment=eq_b, outcome=BookingAttemptOutcome.SUCCESS),
    ]
    urgent = [
        UrgentBookingRequest.objects.create(user=alice, equipment=eq_a, request_type=UrgentBookingRequestType.NO_SLOT),
        UrgentBookingRequest.objects.create(user=bob, equipment=eq_b, request_type=UrgentBookingRequestType.NO_SLOT),
    ]
    WaitlistEntry.objects.create(user=alice, equipment=eq_a)
    WaitlistEntry.objects.create(user=bob, equipment=eq_b)
    return SimpleNamespace(f=f, eq_a=eq_a, eq_b=eq_b, oic_a=oic_a, admin=admin, alice=alice, bob=bob, logs=logs,
                           urgent=urgent)


@pytest.mark.django_db
def test_unknown_report_and_bad_format(world):
    client = world.f.client_for(world.admin)
    assert _export(client, "no-such-report").status_code == 404
    res = _export(client, "booking-attempt-logs", export_format="docx")
    assert res.status_code == 400 and "xlsx, csv or pdf" in res.data["error"]


@pytest.mark.django_db
def test_attempt_log_matches_list_for_admin_and_oic(world):
    for user, expected in ((world.admin, 3), (world.oic_a, 1)):
        client = world.f.client_for(user)
        listed = client.get("/api/booking-attempt-logs/", {"limit": 200}).data
        rows = _csv(_export(client, "booking-attempt-logs"))
        assert len(rows) - 1 == listed["total_count"] == expected
        assert _column(rows, "Email") == [r["user_email"] for r in listed["results"]]
    oic_rows = _csv(_export(world.f.client_for(world.oic_a), "booking-attempt-logs"))
    assert _column(oic_rows, "User") == ["Alice Rao"]


@pytest.mark.django_db
def test_attempt_log_filters_and_formula_guard(world):
    client = world.f.client_for(world.admin)
    rows = _csv(_export(client, "booking-attempt-logs", outcome="FAILED", equipment_id=world.eq_b.pk))
    assert len(rows) == 2
    assert _column(rows, "User") == ["'=HYPERLINK(\"http://x\")"]
    assert _column(rows, "Outcome") == ["Failed"]


@pytest.mark.django_db
def test_attempt_log_forbidden_for_students(world):
    res = _export(world.f.client_for(world.alice), "booking-attempt-logs")
    assert res.status_code == 403


@pytest.mark.django_db
def test_my_booking_attempts_only_own(world):
    rows = _csv(_export(world.f.client_for(world.alice), "my-booking-attempts", outcome="ALL"))
    assert len(rows) - 1 == 2
    assert all("Alpha" in e or "Beta" in e for e in _column(rows, "Equipment"))
    failed = _csv(_export(world.f.client_for(world.alice), "my-booking-attempts"))
    assert len(failed) - 1 == 1


@pytest.mark.django_db
def test_urgent_requests_scoped_like_list(world):
    oic = world.f.client_for(world.oic_a)
    rows = _csv(_export(oic, "urgent-requests"))
    assert _column(rows, "User") == ["Alice Rao"]
    assert len(_csv(_export(world.f.client_for(world.admin), "urgent-requests"))) - 1 == 2
    assert _export(world.f.client_for(world.alice), "urgent-requests").status_code == 403
    mine = _csv(_export(world.f.client_for(world.alice), "my-urgent-requests"))
    assert len(mine) - 1 == 1 and "Alpha XRD" in mine[1][2]


@pytest.mark.django_db
def test_waitlist_xlsx_scoped_with_kpis(world):
    res = _export(world.f.client_for(world.oic_a), "waitlist", export_format="xlsx")
    assert res.status_code == 200
    assert res["X-Export-Row-Count"] == "1"
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Summary", "Waitlisted bookings", "Filters"]
    ws = wb["Waitlisted bookings"]
    headers = [c.value for c in ws[6]]
    user_col = headers.index("User") + 1
    assert ws.cell(row=7, column=user_col).value == "Alice Rao"
    assert ws.cell(row=8, column=1).value is None


@pytest.mark.django_db
def test_pdf_export_is_a_pdf(world):
    res = _export(world.f.client_for(world.admin), "urgent-requests", export_format="pdf")
    assert res.status_code == 200
    assert res["Content-Type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    assert res["Content-Disposition"].startswith('attachment; filename="urgent-booking-requests_')


@pytest.mark.django_db
def test_oic_substitutes_scope(world):
    res = _export(world.f.client_for(world.oic_a), "oic-substitutes", export_format="xlsx")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Substitutes I appointed", "Equipment I substitute for", "Substitution history",
                             "Filters"]
    assert _export(world.f.client_for(world.alice), "oic-substitutes").status_code == 403


@pytest.mark.django_db
def test_reports_statistics_full_workbook_for_admin(world):
    res = _export(world.f.client_for(world.admin), "reports-statistics", export_format="xlsx")
    assert res.status_code == 200, res.content[:300]
    names = load_workbook(io.BytesIO(res.content)).sheetnames
    for sheet in ("Summary", "Booking status", "Revenue by user type", "Revenue by department",
                  "Equipment usage", "Availability", "Slot outcomes", "Ratings", "Filters"):
        assert sheet in names


@pytest.mark.django_db
def test_reports_statistics_student_gets_only_own_booking_statistics(world):
    res = _export(world.f.client_for(world.alice), "reports-statistics", export_format="xlsx")
    assert res.status_code == 200
    assert load_workbook(io.BytesIO(res.content)).sheetnames == ["Summary", "Booking status", "Filters"]
    assert _export(world.f.client_for(world.alice), "equipment-performance").status_code == 403


@pytest.mark.django_db
def test_single_table_of_a_report(world):
    res = _export(world.f.client_for(world.admin), "equipment-performance", export_format="csv",
                  table="revenue_department")
    rows = _csv(res)
    assert rows[0] == ["Department", "Bookings", "Revenue (₹)"]
    bad = _export(world.f.client_for(world.admin), "equipment-performance", table="nope")
    assert bad.status_code == 400
