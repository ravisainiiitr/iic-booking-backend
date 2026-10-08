"""View Booking / My Bookings export with every user input: input columns, sample sets, tables, files, charges."""

import io
import re
from datetime import timedelta

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext

from iic_booking.equipment import booking_list_export
from iic_booking.equipment.models import DynamicInputField
from iic_booking.equipment.models import EquipmentProfileType
from iic_booking.equipment.models import PrintAnalysis

from .test_booking_list_export import _csv_rows
from .test_booking_list_export import _export
from .test_booking_list_export import world  # noqa: F401 - fixture

LIST_URL = "/api/bookings/"


def _field(equipment, key, label, field_type, **kw):
    return DynamicInputField.objects.create(
        equipment=equipment, field_key=key, field_label=label, field_type=field_type, **kw,
    )


def add_rich_inputs(world):
    """Alpha XRD fields of every type; b1 gets two sample sets, tables, comments, a charge breakdown."""
    eq = world.eq_a
    _field(eq, "A", "No. of Samples", "NUMERIC")
    _field(eq, "B", "Sample form:", "RADIO", options=[{"value": "P", "label": "Powder"}, {"value": "F", "label": "Thin film"}])
    _field(eq, "C", "Hazardous", "TOGGLE")
    _field(eq, "D", "Scan modes", "MULTI_SELECT", options=["Fast", "Slow", "Rietveld"])
    _field(eq, "E", "Sample list", "TABLE", options=["Sample name", "Composition"])
    _field(
        eq, "F", "Measurement plan", "TYPED_TABLE",
        table_config={"columns": [
            {"key": "range", "label": "2θ range", "type": "TEXT"},
            {"key": "step", "label": "Step (°)", "type": "NUMERIC"},
            {"key": "spin", "label": "Spinner", "type": "TOGGLE"},
        ]},
    )
    _field(eq, "G", "Sample description", "TEXT")
    b1 = world.b1
    b1.input_values = {
        "A": 3,
        "B": "P",
        "C": True,
        "D": ["Fast", "Rietveld"],
        "E": [["S1", "TiO2"], ["S2", "ZnO"]],
        "F": [{"range": "10-80", "step": 0.02, "spin": True}, {"range": "20-60", "step": 0.01, "spin": False}],
        "G": "नमूना विवरण — dried at 80 °C",
        "comments": "Return the sample after analysis",
        "internal_note": "SECRET-UNDEFINED-KEY",
        "_audit": "SECRET-UNDERSCORE-KEY",
        "_sample_sets": [{"A": 2, "B": "F", "G": "Second set film"}],
    }
    b1.charge_breakdown = [
        {"description": "Analysis charge (3 samples)", "amount": 1271.19},
        {"description": "GST 18%", "amount": 228.81},
    ]
    b1.atmosphere_sensitive_sample = True
    b1.save(update_fields=["input_values", "charge_breakdown", "atmosphere_sensitive_sample"])
    return b1


def add_fabrication_booking(world):
    f = world.f
    printer = f.equipment(name="Gamma 3D Printer", profile_type=EquipmentProfileType.PRINT_3D)
    booking = f.booking(world.alice, printer, world.base + timedelta(days=4), input_values={
        "A": 2, "_quantity_in_a": True,
    })
    for name, cancelled in (("bracket.stl", None), ("removed.stl", world.base)):
        PrintAnalysis.objects.create(
            equipment=printer, user=world.alice, booking=booking, stl_file=f"print_stl/{name}",
            original_filename=name, part_name="Bracket" if name == "bracket.stl" else "", quantity=2,
            material_code_snapshot="PLA", cancelled_at=cancelled,
        )
    return booking


def _by_id(rows):
    header = rows[0]
    return header, {r[1]: dict(zip(header, r)) for r in rows[1:]}


@pytest.mark.django_db
def test_csv_has_one_column_per_input_with_readable_values(world):
    b1 = add_rich_inputs(world)
    header, by_id = _by_id(_csv_rows(_export(world.f.client_for(world.admin), view="staff")))
    row = by_id[b1.virtual_booking_id]
    assert row["No. of Samples (A)"] == "Set 1: 3 | Set 2: 2"
    assert row["Sample form (B)"] == "Set 1: Powder | Set 2: Thin film"
    assert row["Hazardous (C)"] == "Set 1: Yes | Set 2: —"
    assert row["Scan modes (D)"].startswith("Set 1: Fast, Rietveld")
    assert "Sample name: S1, Composition: TiO2; Sample name: S2, Composition: ZnO" in row["Sample list (E)"]
    assert "2θ range: 10-80, Step (°): 0.02, Spinner: Yes" in row["Measurement plan (F)"]
    assert row["Sample description (G)"] == "Set 1: नमूना विवरण — dried at 80 °C | Set 2: Second set film"
    assert row["Any other requirements"] == "Return the sample after analysis"
    assert row["Atmosphere-sensitive sample"].startswith("Yes")
    assert row["Charge breakdown (₹)"] == "Analysis charge (3 samples): 1271.19; GST 18%: 228.81"
    # Inputs are grouped after the list columns, in field order.
    assert header.index("No. of Samples (A)") > header.index("Booked on (IST)")
    assert header.index("Sample form (B)") == header.index("No. of Samples (A)") + 1
    # Bookings without inputs leave the input cells empty.
    assert by_id[world.b2.virtual_booking_id]["No. of Samples (A)"] == ""


@pytest.mark.django_db
def test_keys_without_a_field_definition_and_internal_keys_are_not_exported(world):
    add_rich_inputs(world)
    for fmt in ("csv", "xlsx"):
        content = _export(world.f.client_for(world.admin), fmt, view="staff").content
        if fmt == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content))
            content = repr([list(ws.values) for ws in wb.worksheets]).encode()
        assert b"SECRET-UNDEFINED-KEY" not in content
        assert b"SECRET-UNDERSCORE-KEY" not in content
        assert b"Internal note" not in content


@pytest.mark.django_db
def test_operators_get_inputs_but_no_charges(world):
    from openpyxl import load_workbook

    b1 = add_rich_inputs(world)
    client = world.f.client_for(world.operator_a)
    header, by_id = _by_id(_csv_rows(_export(client, view="staff")))
    assert by_id[b1.virtual_booking_id]["No. of Samples (A)"] == "Set 1: 3 | Set 2: 2"
    assert "Charge breakdown (₹)" not in header and "Amount (₹)" not in header
    wb = load_workbook(io.BytesIO(_export(client, "xlsx", view="staff").content))
    assert wb.sheetnames == ["Bookings", "Sample sets", "Input tables", "Filters"]
    flat = repr([list(ws.values) for ws in wb.worksheets])
    assert "1271.19" not in flat and "GST 18%" not in flat


@pytest.mark.django_db
def test_xlsx_has_sample_sets_input_tables_and_charges_sheets(world):
    from openpyxl import load_workbook

    b1 = add_rich_inputs(world)
    res = _export(world.f.client_for(world.admin), "xlsx", view="staff")
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Bookings", "Sample sets", "Input tables", "Charges", "Filters"]

    ws = wb["Bookings"]
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == f"A1:{ws.cell(row=1, column=ws.max_column).column_letter}{ws.max_row}"
    header = [c.value for c in ws[1]]
    assert "No. of Samples (A)" in header and "Measurement plan (F)" in header
    assert ws.cell(row=1, column=1).fill.fgColor.rgb.endswith("153F79")
    amount = ws.cell(row=2, column=header.index("Amount (₹)") + 1)
    assert "₹" in amount.number_format

    sets = list(wb["Sample sets"].values)
    assert sets[0][:3] == ("Booking ID", "Equipment", "Sample set")
    b1_sets = [r for r in sets[1:] if r[0] == b1.virtual_booking_id]
    assert [r[2] for r in b1_sets] == [1, 2]
    set_header = list(sets[0])
    assert b1_sets[1][set_header.index("Sample form (B)")] == "Thin film"
    assert b1_sets[0][set_header.index("Hazardous (C)")] == "Yes"

    tables = [r for r in wb["Input tables"].values if any(v is not None for v in r)]
    titles = [r[0] for r in tables if r[1] is None and r[0]]
    assert "Alpha XRD — Measurement plan (F)" in titles and "Alpha XRD — Sample list (E)" in titles
    plan_header = next(r for r in tables if r[:3] == ("Booking ID", "Sample set", "Row #") and "2θ range" in r)
    assert list(plan_header[3:6]) == ["2θ range", "Step (°)", "Spinner"]
    assert (b1.virtual_booking_id, 1, 2, "20-60", "0.01", "No") in [tuple(r[:6]) for r in tables]

    charges = list(wb["Charges"].values)
    assert charges[0] == ("Booking ID", "Equipment", "Line", "Description", "Amount (₹)")
    assert (b1.virtual_booking_id, "Alpha XRD", 2, "GST 18%", 228.81) in charges

    info = {r[0]: r[1] for r in wb["Filters"].values if r and r[0]}
    assert info["Exported from"] == "View Booking"
    assert info["Booked"] == 2 and info["Completed"] == 1


@pytest.mark.django_db
def test_fabrication_quantity_and_uploaded_file_names(world):
    booking = add_fabrication_booking(world)
    header, by_id = _by_id(_csv_rows(_export(world.f.client_for(world.admin), view="staff")))
    row = by_id[booking.virtual_booking_id]
    assert row["Quantity Required (A)"] == "2"
    assert row["Uploaded files"] == "bracket.stl (Bracket, PLA, ×2)"
    assert "removed.stl" not in row["Uploaded files"]
    assert "print_stl/" not in repr(by_id)


@pytest.mark.django_db
def test_my_bookings_export_includes_own_inputs_and_charges(world):
    b1 = add_rich_inputs(world)
    header, by_id = _by_id(_csv_rows(_export(world.f.client_for(world.alice), view="my")))
    assert by_id[b1.virtual_booking_id]["Sample list (E)"]
    assert "Charge breakdown (₹)" in header


@pytest.mark.django_db
def test_pdf_has_cover_summary_and_cards_with_embedded_unicode_font(world):
    add_rich_inputs(world)
    add_fabrication_booking(world)
    res = _export(world.f.client_for(world.admin), "pdf", view="staff")
    assert res.status_code == 200, getattr(res, "data", "")
    assert res.content.startswith(b"%PDF")
    assert b"NotoSans" in res.content and b"NotoSansDevanagari" in res.content
    media_box = re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", res.content)
    assert media_box and float(media_box.group(1)) < float(media_box.group(2))  # A4 portrait
    assert len(re.findall(rb"/Type /Page\b", res.content)) >= 2


@pytest.mark.django_db
def test_pdf_handles_values_longer_than_a_page(world):
    b1 = add_rich_inputs(world)
    long_text = "Very long note. " * 600
    b1.input_values = {**b1.input_values, "G": long_text, "comments": long_text,
                       "E": [["x" * 9000, "y"]], "_sample_sets": []}
    b1.save(update_fields=["input_values"])
    res = _export(world.f.client_for(world.admin), "pdf", view="staff")
    assert res.status_code == 200
    assert len(re.findall(rb"/Type /Page\b", res.content)) >= 4


@pytest.mark.django_db
def test_pdf_falls_back_to_helvetica_without_font_files(world, settings, tmp_path, monkeypatch):
    from reportlab.pdfbase import pdfmetrics

    from iic_booking.equipment import export_styles

    add_rich_inputs(world)
    settings.BASE_DIR = tmp_path
    monkeypatch.setattr(pdfmetrics, "getRegisteredFontNames", lambda: [])
    monkeypatch.setattr("iic_booking.equipment.document_exports._register_pdf_rupee_font", lambda: None)
    fonts = export_styles.register_fonts()
    assert fonts == export_styles.Fonts("Helvetica", "Helvetica-Bold", None, False)
    assert export_styles.markup("₹5 – <b>", fonts) == "Rs.5 - &lt;b&gt;"


def test_markup_puts_devanagari_runs_in_the_devanagari_font():
    from iic_booking.equipment.export_styles import Fonts
    from iic_booking.equipment.export_styles import markup

    fonts = Fonts("IICNotoSans", "IICNotoSans-Bold", "IICNotoDeva", True)
    assert markup("Sample नमूना विवरण & ₹10", fonts) == (
        'Sample <font name="IICNotoDeva">नमूना विवरण</font> &amp; ₹10'
    )


@pytest.mark.django_db
def test_pdf_is_capped_lower_than_spreadsheets(world, monkeypatch):
    monkeypatch.setattr(booking_list_export, "PDF_ROW_LIMIT", 2)
    client = world.f.client_for(world.admin)
    res = _export(client, "pdf", view="staff")
    assert res.status_code == 400
    assert "limited to 2" in res.data["error"] and "Excel" in res.data["error"]
    assert _export(client, "xlsx", view="staff").status_code == 200
    assert _export(client, "pdf", view="staff", status="COMPLETED").status_code == 200


@pytest.mark.django_db
@pytest.mark.parametrize("export_format", ["csv", "xlsx", "pdf"])
def test_export_with_inputs_query_count_does_not_grow_per_booking(world, export_format):
    add_rich_inputs(world)
    add_fabrication_booking(world)
    f = world.f
    client = f.client_for(world.admin)

    def count():
        with CaptureQueriesContext(connection) as ctx:
            assert _export(client, export_format, view="staff").status_code == 200
        return len(ctx.captured_queries)

    before = count()
    for i in range(12):
        f.booking(f.student(), world.eq_a, world.base + timedelta(days=8, hours=i % 8),
                  input_values={"A": i + 1, "E": [["x", "y"]], "_sample_sets": [{"A": 1}]})
    assert count() <= before + 2
    assert count() <= 30


@pytest.mark.django_db
def test_list_endpoint_allows_500_rows_per_page_for_the_list_view(world):
    client = world.f.client_for(world.admin)
    res = client.get(LIST_URL, {"list_view": "true", "limit": 500})
    assert res.status_code == 200 and res.data["limit"] == 500
    assert client.get(LIST_URL, {"list_view": "true", "limit": 5000}).data["limit"] == 500
    assert client.get(LIST_URL, {"list_view": "true", "limit": 25, "offset": 25}).data["offset"] == 25
    # Full booking payloads stay at 100 per request.
    assert client.get(LIST_URL, {"limit": 500}).data["limit"] == 100


@pytest.mark.django_db
def test_list_view_query_count_does_not_grow_with_page_size(world):
    f = world.f
    client = f.client_for(world.admin)

    def count():
        with CaptureQueriesContext(connection) as ctx:
            res = client.get(LIST_URL, {"list_view": "true", "limit": 500})
            assert res.status_code == 200
        return len(ctx.captured_queries), res.data["count"]

    before, n_before = count()
    for i in range(20):
        f.booking(f.student(), world.eq_b, world.base + timedelta(days=9, hours=i % 8))
    after, n_after = count()
    assert n_after == n_before + 20
    assert after <= before + 2


def test_field_label_with_key():
    from iic_booking.equipment.booking_export_details import field_label_with_key

    assert field_label_with_key("A", "No. of Samples:") == "No. of Samples (A)"
    assert field_label_with_key("B", "Mode (B)") == "Mode (B)"
    assert field_label_with_key("sample_type", "") == "Sample type"


def test_shared_export_styles_build_a_styled_workbook():
    import io

    from openpyxl import Workbook
    from openpyxl import load_workbook

    from iic_booking.equipment.export_styles import SheetWriter
    from iic_booking.equipment.export_styles import register_xlsx_styles
    from iic_booking.equipment.export_styles import xlsx_text

    wb = Workbook(write_only=True)
    register_xlsx_styles(wb)
    sheet = SheetWriter(wb, "Items")
    sheet.header(["Name", "Amount"])
    sheet.row([xlsx_text("=cmd"), 12.5], ["text", "money"])
    sheet.row([xlsx_text(""), None], ["text", "money"], striped=True)
    sheet.flush(autofilter_columns=2)
    buf = io.BytesIO()
    wb.save(buf)
    ws = load_workbook(io.BytesIO(buf.getvalue()))["Items"]
    assert ws["A2"].value == "'=cmd" and ws["A3"].value is None
    assert ws["A1"].style == "exp_header" and ws["B2"].style == "exp_money" and ws["A3"].style == "exp_text_alt"
    assert ws.auto_filter.ref == "A1:B3"
