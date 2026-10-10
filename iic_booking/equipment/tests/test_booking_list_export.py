"""View Booking / My Bookings export: same rows as the list endpoint, per-role columns, xlsx / csv / pdf files."""

import csv
import io
import re
from datetime import datetime
from datetime import timedelta
from types import SimpleNamespace

import pytest

from iic_booking.equipment import booking_list_export
from iic_booking.equipment.booking_list_export import guard_formula
from iic_booking.equipment.booking_list_export import slot_summary
from iic_booking.equipment.models import BookingStatus
from iic_booking.equipment.models import EquipmentManager
from iic_booking.equipment.models import EquipmentOperator
from iic_booking.equipment.models import WaitlistEntry
from iic_booking.users.models.user_type import UserType
from iic_booking.users.tests.factories import UserFactory

EXPORT_URL = "/api/bookings/export/"


def _staff(user_type, **kw):
    return UserFactory(user_type=user_type, admin_approved=True, **kw)


def _export(client, export_format="csv", **params):
    return client.get(EXPORT_URL, {"export_format": export_format, **params})


def _csv_rows(response):
    assert response.status_code == 200, getattr(response, "data", response.content[:300])
    raw = response.content
    assert raw.startswith("\ufeff".encode())
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))


def _list_ids(client, **params):
    res = client.get("/api/bookings/", {"list_view": "true", "limit": 100, **params})
    assert res.status_code == 200, res.data
    return [row["booking_id"] for row in res.data["bookings"]]


@pytest.fixture
def world(egs_factory):
    f = egs_factory
    eq_a = f.equipment(name="Alpha XRD")
    eq_b = f.equipment(name="Beta TGA")
    oic_a = _staff(UserType.MANAGER, department=f.department)
    operator_a = _staff(UserType.OPERATOR, department=f.department)
    EquipmentManager.objects.create(equipment=eq_a, manager=oic_a)
    EquipmentOperator.objects.create(equipment=eq_a, operator=operator_a, role=EquipmentOperator.Role.PRIMARY)
    admin = _staff(UserType.ADMIN, is_staff=True)
    alice = f.student()
    alice.name = "Alice Rao"
    alice.save(update_fields=["name"])
    bob = f.student()
    bob.name = '=HYPERLINK("http://x")'
    bob.save(update_fields=["name"])
    base = f.future(days=3, hour=9)
    b1 = f.booking(alice, eq_a, base, slot_count=2, total_charge="1500.00")
    b2 = f.booking(bob, eq_b, base + timedelta(days=1), total_charge="250.50")
    b3 = f.booking(alice, eq_b, base + timedelta(days=2))
    type(b3).objects.filter(pk=b3.pk).update(status=BookingStatus.COMPLETED)
    return SimpleNamespace(
        f=f, eq_a=eq_a, eq_b=eq_b, oic_a=oic_a, operator_a=operator_a, admin=admin, alice=alice, bob=bob,
        b1=b1, b2=b2, b3=b3, base=base,
    )


@pytest.mark.django_db
def test_csv_has_bom_header_and_every_matching_booking(world):
    client = world.f.client_for(world.admin)
    rows = _csv_rows(_export(client, view="staff", ordering="-created_at"))
    header = rows[0]
    assert header[:3] == ["S.No.", "Booking ID", "Equipment"]
    assert "Amount (₹)" in header and "Supervisor" in header and "Booked on (IST)" in header
    body = rows[1:]
    assert [r[1] for r in body] == _list_ids(client, ordering="-created_at")
    assert [r[0] for r in body] == [str(i) for i in range(1, len(body) + 1)]
    by_id = {r[1]: dict(zip(header, r)) for r in body}
    first = by_id[world.b1.virtual_booking_id]
    assert first["User"] == "Alice Rao"
    assert first["Amount (₹)"] == "1500.00"
    assert first["Duration"] == "2h"
    day = world.base.strftime("%d-%m-%Y")
    assert first["Slot date(s)"] == day
    assert first["Slot time(s) (IST)"] == "09:00–11:00"
    assert re.fullmatch(r"\d{2}-\d{2}-\d{4} \d{2}:\d{2}", first["Booked on (IST)"])


@pytest.mark.django_db
def test_csv_and_xlsx_guard_formula_cells(world):
    client = world.f.client_for(world.admin)
    rows = _csv_rows(_export(client, view="staff", search=world.b2.virtual_booking_id))
    header = rows[0]
    assert dict(zip(header, rows[1]))["User"] == "'=HYPERLINK(\"http://x\")"

    from openpyxl import load_workbook

    res = _export(client, "xlsx", view="staff", search=world.b2.virtual_booking_id)
    ws = load_workbook(io.BytesIO(res.content))["Bookings"]
    user_col = [c.value for c in ws[1]].index("User") + 1
    cell = ws.cell(row=2, column=user_col)
    assert cell.value == "'=HYPERLINK(\"http://x\")"
    assert cell.data_type == "s"


def test_guard_formula_rules():
    assert guard_formula("=1+1") == "'=1+1"
    assert guard_formula("+91 98765") == "'+91 98765"
    assert guard_formula("-x") == "'-x"
    assert guard_formula("@SUM(A1)") == "'@SUM(A1)"
    assert guard_formula("Alice") == "Alice"
    assert guard_formula(12) == 12


@pytest.mark.django_db
def test_xlsx_bold_frozen_header_and_typed_cells(world):
    from openpyxl import load_workbook

    client = world.f.client_for(world.admin)
    res = _export(client, "xlsx", view="staff", equipment_id=world.eq_a.pk)
    assert res.status_code == 200
    assert res["Content-Type"].startswith("application/vnd.openxmlformats")
    assert re.fullmatch(r'attachment; filename="bookings_\d{4}-\d{2}-\d{2}_\d{4}\.xlsx"', res["Content-Disposition"])
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Bookings"]
    assert ws.freeze_panes == "C2"
    header = [c.value for c in ws[1]]
    assert all(c.font.bold for c in ws[1])
    assert ws.max_row == 2
    row = {h: ws.cell(row=2, column=i + 1) for i, h in enumerate(header)}
    assert row["Booking ID"].value == world.b1.virtual_booking_id
    assert isinstance(row["Booked on (IST)"].value, datetime)
    assert row["Booked on (IST)"].number_format == "DD-MM-YYYY HH:MM"
    assert row["Amount (₹)"].value == 1500.0
    assert ws.column_dimensions["C"].width >= len("Alpha XRD")
    info = {r[0]: r[1] for r in wb["Filters"].iter_rows(min_row=2, values_only=True)}
    assert info["Equipment"] == "Alpha XRD"
    assert info["Rows"] == "1"


@pytest.mark.django_db
def test_pdf_is_a4_portrait_with_a_card_per_booking_for_long_lists(world, monkeypatch):
    f = world.f
    for i in range(70):
        f.booking(world.alice, world.eq_b, world.base + timedelta(days=5, hours=i % 8))
    res = _export(f.client_for(world.admin), "pdf", view="staff")
    assert res.status_code == 200
    assert res["Content-Type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    pages = re.findall(rb"/Type /Page\b", res.content)
    assert len(pages) >= 10
    media_box = re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", res.content)
    assert media_box and float(media_box.group(1)) < float(media_box.group(2))
    assert res["X-Export-Row-Count"] == "73"


@pytest.mark.django_db
@pytest.mark.parametrize(
    "params",
    [
        {},
        {"status": "BOOKED"},
        {"status": "COMPLETED"},
        {"search": "alpha"},
        {"user_name": "alice"},
        {"ordering": "user_name"},
        {"ordering": "-start_time"},
        {"ordering": "booking_ref"},
        {"user_type_filter": "internal"},
        {"start_date": "2000-01-01", "end_date": "2999-12-31", "ordering": "duration"},
    ],
)
def test_export_matches_list_endpoint_for_same_filters(world, params):
    client = world.f.client_for(world.admin)
    rows = _csv_rows(_export(client, view="staff", **params))
    assert [r[1] for r in rows[1:]] == _list_ids(client, **params)


@pytest.mark.django_db
def test_single_character_search_is_ignored_like_the_page(world):
    client = world.f.client_for(world.admin)
    everything = [r[1] for r in _csv_rows(_export(client, view="staff"))[1:]]
    assert [r[1] for r in _csv_rows(_export(client, view="staff", search="q"))[1:]] == everything
    assert [r[1] for r in _csv_rows(_export(client, view="staff", user_name="z"))[1:]] == everything


@pytest.mark.django_db
def test_oic_and_operator_export_only_their_equipment(world):
    for user in (world.oic_a, world.operator_a):
        client = world.f.client_for(user)
        ids = [r[1] for r in _csv_rows(_export(client, view="staff"))[1:]]
        assert ids == [world.b1.virtual_booking_id]
        assert ids == _list_ids(client)
        # Filtering to another OIC's equipment cannot widen the scope.
        assert _csv_rows(_export(client, view="staff", equipment_id=world.eq_b.pk))[1:] == []


@pytest.mark.django_db
def test_amount_column_follows_charge_visibility(world):
    oic_header = _csv_rows(_export(world.f.client_for(world.oic_a), view="staff"))[0]
    operator_header = _csv_rows(_export(world.f.client_for(world.operator_a), view="staff"))[0]
    assert "Amount (₹)" in oic_header
    assert "Amount (₹)" not in operator_header
    assert "Mobile" in operator_header and "Email" in operator_header


@pytest.mark.django_db
def test_users_cannot_use_the_staff_view_and_my_view_is_their_own(world):
    client = world.f.client_for(world.alice)
    assert _export(client, view="staff").status_code == 403
    rows = _csv_rows(_export(client, view="my"))
    header = rows[0]
    assert "Amount (₹)" in header
    assert "Email" not in header and "Mobile" not in header
    assert sorted(r[1] for r in rows[1:]) == sorted([world.b1.virtual_booking_id, world.b3.virtual_booking_id])


@pytest.mark.django_db
def test_my_bookings_lists_active_waitlist_entries_when_status_allows(world):
    f = world.f
    eq_c = f.equipment(name="Gamma SEM")
    WaitlistEntry.objects.create(user=world.alice, equipment=eq_c, status="ACTIVE")
    WaitlistEntry.objects.create(user=world.alice, equipment=world.eq_a, status="OPT_OUT")
    client = f.client_for(world.alice)

    rows = _csv_rows(_export(client, view="my", ordering="-created_at"))
    statuses = {r[2]: r[rows[0].index("Status")] for r in rows[1:]}
    assert statuses["Gamma SEM"].startswith("Waitlisted (WL1")
    assert len(rows) == 4  # header + two bookings + one active waitlist entry

    waitlisted = _csv_rows(_export(client, view="my", status="WAITLISTED"))[1:]
    assert [r[2] for r in waitlisted] == ["Gamma SEM"]
    assert waitlisted[0][1].endswith("W")

    assert all(r[2] != "Gamma SEM" for r in _csv_rows(_export(client, view="my", status="BOOKED"))[1:])
    assert all(r[2] != "Gamma SEM" for r in _csv_rows(_export(client, view="my", equipment_id=world.eq_b.pk))[1:])


@pytest.mark.django_db
def test_finance_my_bookings_export_has_no_waitlist_and_keeps_scope(world):
    finance = _staff(UserType.FINANCE)
    WaitlistEntry.objects.create(user=finance, equipment=world.eq_a, status="ACTIVE")
    client = world.f.client_for(finance)
    rows = _csv_rows(_export(client, view="my"))
    assert [r[1] for r in rows[1:]] == _list_ids(client)
    assert rows[1:] == []


@pytest.mark.django_db
def test_row_cap_returns_a_clear_error(world, monkeypatch):
    monkeypatch.setattr(booking_list_export, "EXPORT_ROW_LIMIT", 2)
    res = _export(world.f.client_for(world.admin), view="staff")
    assert res.status_code == 400
    assert "limited to 2" in res.data["error"]
    ok = _export(world.f.client_for(world.admin), view="staff", status="COMPLETED")
    assert ok.status_code == 200


@pytest.mark.django_db
def test_rejects_unknown_format_and_view(world):
    client = world.f.client_for(world.admin)
    assert client.get(EXPORT_URL, {"export_format": "docx"}).status_code == 400
    assert _export(client, view="everything").status_code == 400


@pytest.mark.django_db
def test_export_query_count_does_not_grow_per_booking(world, django_assert_max_num_queries):
    f = world.f
    for i in range(15):
        student = f.student()
        f.booking(student, world.eq_b, world.base + timedelta(days=6, hours=i % 8))
    with django_assert_max_num_queries(30):
        res = _export(f.client_for(world.admin), view="staff")
    assert res.status_code == 200


def test_slot_summary_merges_adjacent_slots_and_compresses_long_ranges():
    from django.utils import timezone

    start = timezone.make_aware(datetime(2026, 10, 12, 9, 0))
    hour = timedelta(hours=1)
    assert slot_summary([(start, start + hour), (start + hour, start + 2 * hour)]) == ("12-10-2026", "09:00–11:00")
    days = [(start + timedelta(days=d), start + timedelta(days=d) + hour) for d in range(5)]
    assert slot_summary(days) == ("12-10-2026 to 16-10-2026 (5 days)", "09:00–10:00")
    assert slot_summary([]) == ("", "")
