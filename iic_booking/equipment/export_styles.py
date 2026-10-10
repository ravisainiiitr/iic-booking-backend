"""Shared look for portal exports: PDF fonts, colours and paragraph styles; Excel named styles and sheet writer.

PDF text uses the Noto Sans fonts vendored in the repository's ``fonts`` folder (₹ and Devanagari; SIL Open
Font Licence alongside), falling back to DejaVu or Helvetica (with "Rs.") when they are missing. Excel sheets
use the ``exp_*`` named styles from :func:`register_xlsx_styles` and are written with :class:`SheetWriter`.
"""

from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass
from types import SimpleNamespace

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Spreadsheet safety
# ---------------------------------------------------------------------------

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def guard_formula(value):
    """Spreadsheet formula-injection guard: text starting with = + - @ (or tab / CR) gets a leading apostrophe."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

BRAND = "#153f79"
INK = "#1e293b"
MUTED = "#64748b"
RULE = "#cbd5e1"
LABEL_BG = "#f1f5f9"
STRIPE_BG = "#f8fafc"
TABLE_HEAD_BG = "#e3ebf6"
SET_BG = "#eef3fa"

# (text, background) of a booking status badge
STATUS_COLOURS = {
    "BOOKED": ("#1d4ed8", "#dbeafe"),
    "COMPLETED": ("#15803d", "#dcfce7"),
    "PROCESSING": ("#0f766e", "#ccfbf1"),
    "PENDING": ("#b45309", "#fef3c7"),
    "PENDING_PAYMENT": ("#b45309", "#fef3c7"),
    "HOLD": ("#b45309", "#fef3c7"),
    "DISRUPTION_PENDING": ("#c2410c", "#ffedd5"),
    "UNDER_MAINTENANCE": ("#c2410c", "#ffedd5"),
    "OTHER_DISRUPTION": ("#c2410c", "#ffedd5"),
    "ABSENT": ("#c2410c", "#ffedd5"),
    "BOOKING_NOT_UTILIZED": ("#c2410c", "#ffedd5"),
    "CANCELLED": ("#b91c1c", "#fee2e2"),
    "REFUNDED": ("#475569", "#e2e8f0"),
    "WAITLISTED": ("#7e22ce", "#f3e8ff"),
}
DEFAULT_STATUS_COLOURS = ("#475569", "#e2e8f0")

_DEVANAGARI_CHARS = "\u0900-\u097F\uA8E0-\uA8FF\u1CD0-\u1CFF"
_DEVANAGARI_RUN = re.compile(
    f"[{_DEVANAGARI_CHARS}]+(?:[\\s\u200c\u200d]+[{_DEVANAGARI_CHARS}]+)*",
)
_ASCII_FALLBACK = {"₹": "Rs.", "–": "-", "—": "-", "·": "-", "×": "x", "…": "..."}


@dataclass(frozen=True)
class Fonts:
    regular: str
    bold: str
    devanagari: str | None
    unicode: bool


def register_fonts() -> Fonts:
    """Noto Sans (Latin, ₹) and Noto Sans Devanagari from ``BASE_DIR/fonts``; DejaVu or Helvetica otherwise."""
    from django.conf import settings
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dir = os.path.join(str(getattr(settings, "BASE_DIR", "")), "fonts")

    def family(name: str, regular_file: str, bold_file: str) -> bool:
        registered = set(pdfmetrics.getRegisteredFontNames())
        try:
            for font_name, filename in ((name, regular_file), (f"{name}-Bold", bold_file)):
                if font_name not in registered:
                    pdfmetrics.registerFont(TTFont(font_name, os.path.join(font_dir, filename)))
        except Exception:  # noqa: BLE001 - missing or unreadable font file
            logger.warning("PDF export: font %s not available in %s", regular_file, font_dir)
            return False
        pdfmetrics.registerFontFamily(
            name, normal=name, bold=f"{name}-Bold", italic=name, boldItalic=f"{name}-Bold",
        )
        return True

    if family("IICNotoSans", "NotoSans-Regular.ttf", "NotoSans-Bold.ttf"):
        deva = "IICNotoDeva" if family(
            "IICNotoDeva", "NotoSansDevanagari-Regular.ttf", "NotoSansDevanagari-Bold.ttf",
        ) else None
        return Fonts("IICNotoSans", "IICNotoSans-Bold", deva, True)

    from .document_exports import _register_pdf_rupee_font

    dejavu = _register_pdf_rupee_font()
    if dejavu:
        return Fonts(dejavu, dejavu, None, True)
    return Fonts("Helvetica", "Helvetica-Bold", None, False)


def markup(text, fonts: Fonts) -> str:
    """Paragraph markup for plain text: escaped, Devanagari runs in the Devanagari font, newlines as breaks."""
    text = "" if text is None else str(text)
    if not fonts.unicode:
        for char, replacement in _ASCII_FALLBACK.items():
            text = text.replace(char, replacement)
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    if fonts.devanagari:
        text = _DEVANAGARI_RUN.sub(lambda m: f'<font name="{fonts.devanagari}">{m.group(0)}</font>', text)
    return text.replace("\n", "<br/>")


LINK = "#1d4ed8"
# Table cells are centred like the portal's tables; longer free text reads better left-aligned.
LONG_CELL_TEXT = 60


def link_markup(text_markup: str, url: str, *, color: str = LINK) -> str:
    """Wrap Paragraph markup in a clickable, underlined link (``text_markup`` is already escaped)."""
    if not url:
        return text_markup
    href = url.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<link href="{href}" color="{color}"><u>{text_markup}</u></link>'


def money_text(amount, fonts: Fonts) -> str:
    if amount is None:
        return ""
    return f"{'₹' if fonts.unicode else 'Rs. '}{amount:,.2f}"


def pdf_styles(fonts: Fonts) -> SimpleNamespace:
    """Paragraph styles of the export reports (title, headings, key/value labels, table cells, badges)."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    def style(name, *, bold=False, size=8.0, leading=None, color=INK, **kw):
        return ParagraphStyle(
            name, fontName=fonts.bold if bold else fonts.regular, fontSize=size,
            leading=leading or round(size * 1.3, 1), textColor=colors.HexColor(color), **kw,
        )

    return SimpleNamespace(
        title=style("exp_title", bold=True, size=22, leading=27, color=BRAND, alignment=TA_CENTER),
        subtitle=style("exp_subtitle", size=10, color=MUTED, alignment=TA_CENTER),
        h2=style("exp_h2", bold=True, size=12.5, leading=16, color=BRAND, spaceBefore=4, spaceAfter=5),
        panel_title=style("exp_panel_title", bold=True, size=9, color=BRAND),
        section=style("exp_section", bold=True, size=9.2, leading=12, color=BRAND),
        label=style("exp_label", size=7.3, leading=9.2, color=MUTED),
        value=style("exp_value", size=8, leading=10.2),
        cell=style("exp_cell", size=7.4, leading=9.3),
        cell_small=style("exp_cell_small", size=6.6, leading=8.3),
        cell_center=style("exp_cell_center", size=7.4, leading=9.3, alignment=TA_CENTER),
        cell_small_center=style("exp_cell_small_center", size=6.6, leading=8.3, alignment=TA_CENTER),
        cell_right=style("exp_cell_right", size=7.4, leading=9.3, alignment=TA_RIGHT),
        cell_bold_right=style("exp_cell_bold_right", bold=True, size=7.6, leading=9.5, alignment=TA_RIGHT),
        cell_bold=style("exp_cell_bold", bold=True, size=7.6, leading=9.5),
        head=style("exp_head", bold=True, size=7.4, leading=9.3, color="#ffffff", alignment=TA_CENTER),
        table_head=style("exp_table_head", bold=True, size=7.2, leading=9, color=BRAND, alignment=TA_CENTER),
        table_head_small=style("exp_table_head_small", bold=True, size=6.5, leading=8.2, color=BRAND,
                               alignment=TA_CENTER),
        table_label=style("exp_table_label", bold=True, size=7.6, leading=9.6, color="#334155", spaceBefore=3,
                          spaceAfter=2, keepWithNext=1),
        band_id=style("exp_band_id", bold=True, size=11.5, leading=14, color="#ffffff"),
        band_sub=style("exp_band_sub", size=8.5, leading=11, color="#dbe7f7"),
        badge=style("exp_badge", bold=True, size=7.4, leading=9, alignment=TA_CENTER),
        set_title=style("exp_set_title", bold=True, size=8.4, leading=10.5, color=BRAND),
        note=style("exp_note", size=7.6, leading=9.8, color=MUTED),
        count=style("exp_count", bold=True, size=8, leading=10, alignment=TA_RIGHT),
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

XLSX_BRAND = "153F79"
XLSX_TEXT_MAX = 32_000  # Excel holds at most 32,767 characters in a cell


def xlsx_text(value):
    """Cell text with the formula guard, shortened to fit a cell; empty text becomes an empty cell."""
    text = guard_formula(value)
    if isinstance(text, str) and len(text) > XLSX_TEXT_MAX:
        text = text[:XLSX_TEXT_MAX] + "…"
    return None if text == "" else text


def register_xlsx_styles(wb) -> None:
    """Named styles: exp_text / exp_longtext / exp_int / exp_money / exp_datetime / exp_link (and *_alt stripes),
    exp_header, exp_band, exp_title, exp_label and exp_note. Cells are centred like the portal's tables;
    ``longtext`` (free text such as comments) is left-aligned."""
    from openpyxl.styles import Alignment
    from openpyxl.styles import Border
    from openpyxl.styles import Font
    from openpyxl.styles import NamedStyle
    from openpyxl.styles import PatternFill
    from openpyxl.styles import Side

    def fill(color):
        return PatternFill("solid", start_color=color, end_color=color)

    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body = Font(name="Calibri", size=10, color="1E293B")
    link = Font(name="Calibri", size=10, color=LINK.lstrip("#").upper(), underline="single")
    top_left = Alignment(vertical="top", wrap_text=True)
    centred = Alignment(vertical="center", horizontal="center", wrap_text=True)
    kinds = {
        "text": {"alignment": centred},
        "longtext": {"alignment": Alignment(vertical="center", horizontal="left", wrap_text=True)},
        "int": {"alignment": centred},
        "money": {"alignment": centred, "number_format": '"₹"#,##0.00'},
        "datetime": {"alignment": centred, "number_format": "DD-MM-YYYY HH:MM"},
        "link": {"alignment": centred, "font": link},
    }
    for kind, extra in kinds.items():
        extra = {"font": body, **extra}
        wb.add_named_style(NamedStyle(name=f"exp_{kind}", border=border, **extra))
        wb.add_named_style(NamedStyle(name=f"exp_{kind}_alt", border=border, fill=fill("F4F7FB"), **extra))
    wb.add_named_style(NamedStyle(
        name="exp_header", font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=fill(XLSX_BRAND),
        border=border, alignment=Alignment(vertical="center", horizontal="center", wrap_text=True),
    ))
    wb.add_named_style(NamedStyle(
        name="exp_band", font=Font(name="Calibri", size=11, bold=True, color=XLSX_BRAND), fill=fill("DBE7F7"),
    ))
    wb.add_named_style(NamedStyle(name="exp_title", font=Font(name="Calibri", size=14, bold=True, color=XLSX_BRAND)))
    wb.add_named_style(NamedStyle(
        name="exp_label", font=Font(name="Calibri", size=10, bold=True, color="334155"), fill=fill("EEF2F7"),
        border=border, alignment=top_left,
    ))
    wb.add_named_style(NamedStyle(name="exp_note", font=Font(name="Calibri", size=10, italic=True, color="64748B")))


class SheetWriter:
    """Buffers a write-only sheet so column widths (written before the rows) fit the content.

    Cell kinds are "text", "longtext", "int", "money" and "datetime"; a cell with a link is underlined and
    clickable. When the sheet starts with its header row, that row stays frozen. The workbook needs
    :func:`register_xlsx_styles`.
    """

    def __init__(self, wb, title: str, *, max_width: int = 50):
        self.ws = wb.create_sheet(title)
        self.max_width = max_width
        self.rows: list[list[tuple]] = []
        self.widths: dict[int, int] = {}
        self.data_rows = 0
        self.links: dict[tuple[int, int], str] = {}

    def _track(self, index: int, value, kind: str, *, header: bool = False) -> None:
        if value is None:
            return
        if kind == "datetime":
            length = 16
        elif kind == "money" and not isinstance(value, str):
            length = len(f"₹{value:,.2f}")
        elif header:
            words = str(value).split()
            length = max([min(len(str(value)), 22), *(len(w) for w in words)])
        else:
            length = max((len(line) for line in str(value).split("\n")), default=0)
        self.widths[index] = max(self.widths.get(index, 0), length)

    def header(self, labels) -> None:
        for i, label in enumerate(labels):
            self._track(i, label, "text", header=True)
        self.rows.append([(label, "exp_header") for label in labels])

    def row(self, values, kinds, *, striped: bool = False, links=None) -> None:
        """``links``: optional URL per column (None for plain cells)."""
        out = []
        links = list(links or [])
        for i, (value, kind) in enumerate(zip(values, kinds)):
            self._track(i, value, kind)
            url = links[i] if i < len(links) else None
            if url and value not in (None, ""):
                self.links[(len(self.rows), i)] = url
                kind = "link"
            out.append((value, f"exp_{kind}_alt" if striped else f"exp_{kind}"))
        self.rows.append(out)
        self.data_rows += 1

    def line(self, value, style: str) -> None:
        self.rows.append([(value, style)])

    def blank(self) -> None:
        self.rows.append([])

    def flush(self, *, autofilter_columns: int = 0, min_width: int = 8) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter

        for index, width in self.widths.items():
            letter = get_column_letter(index + 1)
            self.ws.column_dimensions[letter].width = min(max(width + 2, min_width), self.max_width)
        if autofilter_columns and self.data_rows:
            last = get_column_letter(autofilter_columns)
            self.ws.auto_filter.ref = f"A1:{last}{self.data_rows + 1}"
        if self.ws.freeze_panes is None and self.rows and self.rows[0] and self.rows[0][0][1] == "exp_header":
            self.ws.freeze_panes = "A2"
        for row_index, row in enumerate(self.rows):
            cells = []
            for col_index, (value, style) in enumerate(row):
                cell = WriteOnlyCell(self.ws, value=value)
                cell.style = style
                url = self.links.get((row_index, col_index))
                if url:
                    cell.hyperlink = url
                cells.append(cell)
            self.ws.append(cells)
