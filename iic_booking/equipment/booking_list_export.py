"""Download the View Booking / My Bookings list as xlsx, csv or pdf.

Rows are every booking the page would list for the same filters, search, sort and role scope
(the caller passes the ``_booking_list_queryset`` result). Columns follow what each page shows:
``staff`` is View Booking (/booking-management), ``my`` is My Bookings. Each booking also carries every
user input shown in its booking details (``booking_export_details``): one column per input field in csv /
xlsx (plus Sample sets, Input tables and Charges sheets in xlsx) and a card per booking in the pdf.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from decimal import InvalidOperation
from zoneinfo import ZoneInfo

from django.db.models import Prefetch
from django.http import HttpResponse
from django.utils import timezone

from iic_booking.communication.utils import booking_display_id_for_email
from iic_booking.users.display import get_user_display_name
from iic_booking.users.models.user_type import UserType

from .models import Booking
from .models import BookingSampleTrace
from .models import BookingSlotRange
from .models import BookingStatus
from .models import DailySlot
from .models import WaitlistEntry

EXPORT_ROW_LIMIT = 10_000
# The pdf has a details card (often a page) per booking; larger lists are for Excel.
PDF_ROW_LIMIT = 1_000
IST = ZoneInfo("Asia/Kolkata")
PORTAL_HEADER = "Institute Instrumentation Centre (IIC), IIT Roorkee"

# Lab Operators see the job sheet (no charges) on View Booking; OIC / Department Admin / Admin see Total Cost.
STAFF_AMOUNT_USER_TYPES = frozenset({UserType.MANAGER, UserType.DEPT_ADMIN, UserType.ADMIN})

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

_SORT_LABELS = {
    "booking_ref": "Booking ID",
    "booking_id": "Booking ID",
    "equipment_name": "Equipment",
    "user_name": "User name",
    "supervisor_name": "Supervisor name",
    "user_phone": "User mobile",
    "user_email": "User email",
    "start_time": "Start",
    "end_time": "End",
    "duration": "Duration",
    "total_time_minutes": "Duration",
    "total_charge": "Cost",
    "status": "Status",
    "created_at": "Booked on",
    "updated_at": "Last updated",
    "rating": "Rating",
}


class ExportTooLarge(Exception):
    pass


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    pdf_width: float  # relative share of the PDF table width
    kind: str = "text"  # text | int | money | datetime


def charges_visible(view: str, user) -> bool:
    """My Bookings shows the user's own charges; on View Booking Lab Operators see no charges."""
    return view == "my" or getattr(user, "user_type", None) in STAFF_AMOUNT_USER_TYPES


def export_columns(view: str, user) -> list[Column]:
    if view == "staff":
        cols = [
            Column("sno", "S.No", 0.6, "int"),
            Column("booking_id", "Booking ID", 2.3),
            Column("equipment", "Equipment", 1.9),
            Column("user", "User", 1.5),
            Column("user_type", "User type", 1.0),
            Column("department", "Department", 1.3),
            Column("supervisor", "Supervisor", 1.4),
            Column("mobile", "Mobile", 1.6),
            Column("email", "Email", 2.2),
            Column("slot_dates", "Slot date(s)", 1.5),
            Column("slot_times", "Slot time(s) (IST)", 1.5),
            Column("duration", "Duration", 0.8),
            Column("status", "Status", 1.3),
            Column("samples", "Samples", 0.9),
            Column("sample_status", "Sample status", 1.1),
        ]
        if charges_visible(view, user):
            cols.append(Column("amount", "Amount (₹)", 1.0, "money"))
        cols.append(Column("booked_on", "Booked on (IST)", 1.3, "datetime"))
        return cols
    return [
        Column("sno", "S.No", 0.6, "int"),
        Column("booking_id", "Booking ID", 2.2),
        Column("equipment", "Equipment", 2.6),
        Column("user", "User", 1.8),
        Column("slot_dates", "Slot date(s)", 1.8),
        Column("slot_times", "Slot time(s) (IST)", 1.8),
        Column("duration", "Duration", 1.0),
        Column("amount", "Amount (₹)", 1.1, "money"),
        Column("status", "Status", 2.0),
        Column("samples", "Samples", 1.1),
        Column("booked_on", "Booked on (IST)", 1.5, "datetime"),
    ]


# ---------------------------------------------------------------------------
# Row building
# ---------------------------------------------------------------------------


def _ist(dt):
    if dt is None:
        return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt)
    return dt.astimezone(IST)


def format_ist_date(dt) -> str:
    local = _ist(dt)
    return local.strftime("%d-%m-%Y") if local else ""


def format_ist_datetime(dt) -> str:
    local = _ist(dt)
    return local.strftime("%d-%m-%Y %H:%M") if local else ""


def format_duration(minutes) -> str:
    try:
        total = int(minutes or 0)
    except (TypeError, ValueError):
        return ""
    if total <= 0:
        return ""
    if total < 60:
        return f"{total} min"
    hours, mins = divmod(total, 60)
    return f"{hours}h {mins}m" if mins else f"{hours}h"


def slot_summary(ranges) -> tuple[str, str]:
    """(dates, times) text for a booking's slot ranges, in IST; adjacent slots on a day merge into one range."""
    local = sorted((_ist(s), _ist(e)) for s, e in ranges if s is not None and e is not None)
    if not local:
        return "", ""
    merged: list[list] = []
    for start, end in local:
        if merged and merged[-1][1] == start and merged[-1][0].date() == start.date():
            merged[-1][1] = end
        else:
            merged.append([start, end])
    dates: list[str] = []
    times: list[str] = []
    for start, end in merged:
        d = start.strftime("%d-%m-%Y")
        if d not in dates:
            dates.append(d)
        t = f"{start:%H:%M}–{end:%H:%M}"
        if t not in times:
            times.append(t)
    date_text = ", ".join(dates) if len(dates) <= 3 else f"{dates[0]} to {dates[-1]} ({len(dates)} days)"
    time_text = ", ".join(times[:4]) + (f" (+{len(times) - 4} more)" if len(times) > 4 else "")
    return date_text, time_text


def _money(value):
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _load_bookings(queryset) -> list:
    return list(
        queryset.select_related("user", "user__department", "equipment", "charge_profile").prefetch_related(
            Prefetch(
                "daily_slots",
                queryset=DailySlot.objects.only("id", "booking_id", "start_datetime", "end_datetime").order_by(
                    "start_datetime",
                ),
            ),
            Prefetch(
                "sample_trace_events",
                queryset=BookingSampleTrace.objects.only("id", "booking_id", "status", "created_at").order_by(
                    "created_at", "id",
                ),
            ),
        ),
    )


def build_booking_rows(bookings) -> list[dict]:
    from .booking_sample_summary import SampleCountFieldIndex
    from .booking_sample_summary import booking_sample_summary
    from .serializers import _booking_status_display
    from .serializers import wallet_owner_display_names

    user_type_labels = dict(UserType.get_choices())
    users = {b.user.pk: b.user for b in bookings if b.user_id and b.user}
    supervisors = wallet_owner_display_names(list(users.values()))
    no_slot_ids = [b.pk for b in bookings if not b.daily_slots.all()]
    released = {}
    for start_idx in range(0, len(no_slot_ids), 500):
        chunk = no_slot_ids[start_idx:start_idx + 500]
        for booking_id, start, end in BookingSlotRange.objects.filter(booking_id__in=chunk).values_list(
            "booking_id", "start_datetime", "end_datetime",
        ):
            released[booking_id] = (start, end)
    sample_index = SampleCountFieldIndex()
    sample_index.preload({b.equipment_id for b in bookings})

    rows = []
    for b in bookings:
        slots = [(s.start_datetime, s.end_datetime) for s in b.daily_slots.all()]
        if not slots and b.pk in released:
            slots = [released[b.pk]]
        dates, times = slot_summary(slots)
        summary = booking_sample_summary(b, sample_index) or {}
        sample_parts = []
        if (summary.get("sets") or 0) > 1:
            sample_parts.append(f"{summary['sets']} sets")
        if summary.get("samples") is not None:
            n = summary["samples"]
            n_text = int(n) if float(n).is_integer() else n
            sample_parts.append(f"{n_text} {'sample' if n == 1 else 'samples'}")
        events = list(b.sample_trace_events.all())
        user = b.user if b.user_id else None
        code = (b.user_type_snapshot or "").strip()
        rows.append(
            {
                "pk": b.pk,
                "status_code": b.status,
                "booking_id": booking_display_id_for_email(b),
                "equipment": getattr(b.equipment, "name", "") or "",
                "user": get_user_display_name(user) if user else "",
                "user_type": user_type_labels.get(code, code),
                "department": getattr(getattr(user, "department", None), "name", "") or "",
                "supervisor": supervisors.get(getattr(user, "pk", None)) or "",
                "mobile": getattr(user, "phone_number", "") or "",
                "email": getattr(user, "email", "") or "",
                "slot_dates": dates,
                "slot_times": times,
                "duration": format_duration(b.total_time_minutes),
                "status": _booking_status_display(b),
                "samples": " · ".join(sample_parts),
                "sample_status": events[-1].get_status_display() if events else "",
                "amount": _money(b.total_charge),
                "booked_on": b.created_at,
            },
        )
    return rows


def _waitlist_entries(request, *, include: bool):
    """Active waitlist entries My Bookings lists with the bookings, narrowed by the same filters."""
    if not include:
        return []
    params = request.query_params
    if (params.get("start_date") or "").strip() or (params.get("end_date") or "").strip():
        return []
    if str(params.get("results_overdue") or "").strip().lower() in ("1", "true", "yes"):
        return []
    qs = WaitlistEntry.objects.filter(user=request.user, status="ACTIVE").select_related(
        "equipment", "equipment__internal_department", "user",
    )
    equipment_id = (params.get("equipment_id") or "").strip()
    if equipment_id:
        try:
            qs = qs.filter(equipment_id=int(equipment_id))
        except (TypeError, ValueError):
            return []
    return list(qs.order_by("-created_at"))


def build_waitlist_rows(entries, search: str) -> list[dict]:
    from .waitlist import active_waitlist_position
    from .waitlist import waitlist_virtual_booking_id

    rows = []
    for entry in entries:
        position = active_waitlist_position(entry) or 0
        equipment = entry.equipment
        display_id = waitlist_virtual_booking_id(
            getattr(equipment, "code", "") or "",
            position,
            department_code=Booking.department_code_for_virtual_id(equipment),
            created_at=entry.created_at,
        )
        name = getattr(equipment, "name", "") or ""
        if search:
            needle = search.lower()
            if needle not in name.lower() and needle not in display_id.lower():
                continue
        queue = f"WL{position}" if position else "queue"
        rows.append(
            {
                "pk": None,
                "status_code": BookingStatus.WAITLISTED,
                "booking_id": display_id,
                "equipment": name,
                "user": get_user_display_name(entry.user),
                "user_type": "",
                "department": "",
                "supervisor": "",
                "mobile": "",
                "email": "",
                "slot_dates": "",
                "slot_times": "",
                "duration": "",
                "status": f"Waitlisted ({queue}, not a confirmed booking)",
                "samples": "",
                "sample_status": "",
                "amount": None,
                "booked_on": entry.created_at,
            },
        )
    return rows


# ---------------------------------------------------------------------------
# Filters summary / file naming
# ---------------------------------------------------------------------------


def filters_summary(request, *, view: str) -> list[tuple[str, str]]:
    from .models import Equipment

    params = request.query_params

    def text(name):
        value = (params.get(name) or "").strip()
        return value if len(value) >= 2 else ""

    out: list[tuple[str, str]] = []
    status_value = (params.get("status") or "").strip().upper()
    if str(params.get("results_overdue") or "").strip().lower() in ("1", "true", "yes"):
        out.append(("Status", "Results overdue"))
    elif status_value:
        out.append(("Status", dict(BookingStatus.choices).get(status_value, status_value)))
    else:
        out.append(("Status", "All"))
    if text("search"):
        out.append(("Search", text("search")))
    if (params.get("start_date") or "").strip():
        out.append(("From", _iso_to_dmy(params.get("start_date"))))
    if (params.get("end_date") or "").strip():
        out.append(("To", _iso_to_dmy(params.get("end_date"))))
    equipment_id = (params.get("equipment_id") or "").strip()
    if equipment_id.isdigit():
        name = Equipment.objects.filter(pk=int(equipment_id)).values_list("name", flat=True).first()
        out.append(("Equipment", name or equipment_id))
    if text("user_name"):
        out.append(("User name", text("user_name")))
    if text("supervisor_name"):
        out.append(("Supervisor name", text("supervisor_name")))
    user_type = (params.get("user_type_filter") or "").strip().lower()
    if user_type in ("internal", "external"):
        out.append(("User type", "Internal (students / faculty)" if user_type == "internal" else "External"))
    istem = (params.get("istem_fbr") or "").strip().lower()
    if istem in ("verified", "unverified"):
        out.append(("I-STEM FBR", istem.capitalize()))
    ordering = (params.get("ordering") or "").strip()
    key = ordering.lstrip("-")
    if key in _SORT_LABELS:
        out.append(("Sorted by", f"{_SORT_LABELS[key]} ({'descending' if ordering.startswith('-') else 'ascending'})"))
    return out


def status_counts(rows) -> list[tuple[str, int]]:
    """[(status label, bookings)] most common first."""
    labels = dict(BookingStatus.choices)
    counts: dict[str, int] = {}
    for row in rows:
        code = row.get("status_code") or ""
        label = str(labels.get(code, code or "Unknown"))
        counts[label] = counts.get(label, 0) + 1
    return sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))


def _iso_to_dmy(value) -> str:
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").strftime("%d-%m-%Y")
    except ValueError:
        return str(value)


def export_filename(ext: str, now=None) -> str:
    local = _ist(now or timezone.now())
    return f"bookings_{local:%Y-%m-%d_%H%M}.{ext}"


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------


def guard_formula(value):
    """Spreadsheet formula-injection guard: text starting with = + - @ (or tab / CR) gets a leading apostrophe."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _cell_text(col: Column, value) -> str:
    if value is None or value == "":
        return ""
    if col.kind == "money":
        return f"{value:.2f}"
    if col.kind == "datetime":
        return format_ist_datetime(value)
    return str(value)


INPUT_KEY_PREFIX = "input:"


def attach_details(rows, details_by_pk) -> None:
    """Put each booking's BookingDetail on its row, plus one-line texts for the input columns."""
    from .booking_export_details import field_text
    from .booking_export_details import files_text

    for row in rows:
        detail = details_by_pk.get(row.get("pk"))
        row["detail"] = detail
        if detail is None:
            continue
        for item in detail.fields:
            row[INPUT_KEY_PREFIX + item.label] = field_text(item, detail.sets)
        row["atmosphere"] = "Yes (submit at slot start)" if detail.atmosphere_sensitive else "No"
        row["comments"] = detail.comments
        row["files"] = files_text(detail.files)
        row["charge_lines"] = "; ".join(
            f"{desc or 'Charge'}: {'' if amount is None else f'{amount:.2f}'}".rstrip(": ")
            for desc, amount in detail.charges
        )


def input_field_labels(rows) -> list[str]:
    """Distinct input field labels of the exported bookings, grouped by equipment (A-Z) in field order."""
    by_equipment: dict[str, list[str]] = {}
    for row in rows:
        detail = row.get("detail")
        if detail is None:
            continue
        labels = by_equipment.setdefault(row.get("equipment") or "", [])
        for item in detail.fields:
            if item.label not in labels:
                labels.append(item.label)
    out: list[str] = []
    for equipment in sorted(by_equipment, key=str.lower):
        out.extend(label for label in by_equipment[equipment] if label not in out)
    return out


def detail_columns(rows, *, charges: bool, include_charge_lines: bool) -> list[Column]:
    """Columns after the list columns: one per input field, then the booking-level extras that have values."""
    cols = [Column(INPUT_KEY_PREFIX + label, label, 1.0) for label in input_field_labels(rows)]
    details = [r["detail"] for r in rows if r.get("detail") is not None]
    if any(d.atmosphere_sensitive for d in details):
        cols.append(Column("atmosphere", "Atmosphere-sensitive sample", 1.0))
    if any(d.comments for d in details):
        cols.append(Column("comments", "Any other requirements", 1.0))
    if any(d.files for d in details):
        cols.append(Column("files", "Uploaded files", 1.0))
    if charges and include_charge_lines and any(d.charges for d in details):
        cols.append(Column("charge_lines", "Charge breakdown (₹)", 1.0))
    return cols


def render_csv(columns, rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([c.label for c in columns])
    for row in rows:
        writer.writerow(
            [_cell_text(c, row.get(c.key)) if c.kind != "text" else guard_formula(_cell_text(c, row.get(c.key)))
             for c in columns],
        )
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


_XLSX_BRAND = "153F79"
_XLSX_TEXT_MAX = 32_000  # Excel holds at most 32,767 characters in a cell


def _xlsx_text(value):
    text = guard_formula(value)
    if isinstance(text, str) and len(text) > _XLSX_TEXT_MAX:
        text = text[:_XLSX_TEXT_MAX] + "…"
    return None if text == "" else text


def _register_xlsx_styles(wb) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.styles import Border
    from openpyxl.styles import Font
    from openpyxl.styles import NamedStyle
    from openpyxl.styles import PatternFill
    from openpyxl.styles import Side

    def fill(color):
        return PatternFill("solid", start_color=color, end_color=color)

    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body = Font(name="Calibri", size=10, color="1E293B")
    top_left = Alignment(vertical="top", wrap_text=True)
    kinds = {
        "text": {"alignment": top_left},
        "int": {"alignment": Alignment(vertical="top", horizontal="center")},
        "money": {"alignment": Alignment(vertical="top", horizontal="right"), "number_format": '"₹"#,##0.00'},
        "datetime": {"alignment": Alignment(vertical="top", horizontal="left"), "number_format": "DD-MM-YYYY HH:MM"},
    }
    for kind, extra in kinds.items():
        wb.add_named_style(NamedStyle(name=f"exp_{kind}", font=body, border=border, **extra))
        wb.add_named_style(NamedStyle(name=f"exp_{kind}_alt", font=body, border=border, fill=fill("F4F7FB"), **extra))
    wb.add_named_style(NamedStyle(
        name="exp_header", font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=fill(_XLSX_BRAND),
        border=border, alignment=Alignment(vertical="center", wrap_text=True),
    ))
    wb.add_named_style(NamedStyle(
        name="exp_band", font=Font(name="Calibri", size=11, bold=True, color=_XLSX_BRAND), fill=fill("DBE7F7"),
    ))
    wb.add_named_style(NamedStyle(name="exp_title", font=Font(name="Calibri", size=14, bold=True, color=_XLSX_BRAND)))
    wb.add_named_style(NamedStyle(
        name="exp_label", font=Font(name="Calibri", size=10, bold=True, color="334155"), fill=fill("EEF2F7"),
        border=border, alignment=top_left,
    ))
    wb.add_named_style(NamedStyle(name="exp_note", font=Font(name="Calibri", size=10, italic=True, color="64748B")))


class _SheetWriter:
    """Buffers a write-only sheet so column widths (written before the rows) fit the content."""

    def __init__(self, wb, title: str, *, max_width: int = 50):
        self.ws = wb.create_sheet(title)
        self.max_width = max_width
        self.rows: list[list[tuple]] = []
        self.widths: dict[int, int] = {}
        self.data_rows = 0

    def _track(self, index: int, value, kind: str, *, header: bool = False) -> None:
        if value is None:
            return
        if kind == "datetime":
            length = 16
        elif kind == "money":
            length = len(f"₹{value:,.2f}")
        elif header:
            words = str(value).split()
            length = max([min(len(str(value)), 22), *(len(w) for w in words)])
        else:
            length = max((len(line) for line in str(value).split("\n")), default=0)
        self.widths[index] = max(self.widths.get(index, 0), length)

    def header(self, labels) -> None:
        for i, label in enumerate(labels):
            self._track(i, label, "text", header=True)
        self.rows.append([(label, "exp_header") for label in labels])

    def row(self, values, kinds, *, striped: bool = False) -> None:
        out = []
        for i, (value, kind) in enumerate(zip(values, kinds)):
            self._track(i, value, kind)
            out.append((value, f"exp_{kind}_alt" if striped else f"exp_{kind}"))
        self.rows.append(out)
        self.data_rows += 1

    def line(self, value, style: str) -> None:
        self.rows.append([(value, style)])

    def blank(self) -> None:
        self.rows.append([])

    def flush(self, *, autofilter_columns: int = 0, min_width: int = 8) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter

        for index, width in self.widths.items():
            letter = get_column_letter(index + 1)
            self.ws.column_dimensions[letter].width = min(max(width + 2, min_width), self.max_width)
        if autofilter_columns and self.data_rows:
            last = get_column_letter(autofilter_columns)
            self.ws.auto_filter.ref = f"A1:{last}{self.data_rows + 1}"
        for row in self.rows:
            cells = []
            for value, style in row:
                cell = WriteOnlyCell(self.ws, value=value)
                cell.style = style
                cells.append(cell)
            self.ws.append(cells)


def _xlsx_value(col: Column, value):
    if col.kind == "datetime":
        return _ist(value).replace(tzinfo=None) if value is not None else None
    if col.kind == "money":
        return float(value) if value is not None else None
    if col.kind == "int":
        return value
    return _xlsx_text(_cell_text(col, value))


def _sample_set_rows(rows, labels):
    from .input_display import formatted_as_text

    for row in rows:
        detail = row.get("detail")
        if detail is None or not detail.fields:
            continue
        by_label = {item.label: item for item in detail.fields}
        for index in range(detail.sets):
            values = []
            for label in labels:
                item = by_label.get(label)
                value = item.values[index] if item is not None and index < len(item.values) else None
                values.append(_xlsx_text(formatted_as_text(value)) if value else None)
            yield row, index + 1, values


def _input_table_groups(rows) -> list[tuple[str, list[str], list[list]]]:
    """[(title, columns, rows)] per equipment / table field / column layout; rows are [booking, set, row #, cells…]."""
    groups: dict[tuple, list[list]] = {}
    for row in rows:
        detail = row.get("detail")
        if detail is None:
            continue
        for item in detail.fields:
            for set_index, value in enumerate(item.values, start=1):
                if value.get("kind") != "table":
                    continue
                columns = list(value.get("columns") or [])
                data = [list(r) for r in value.get("rows") or []]
                if columns and columns[0] == "S.No.":
                    columns = columns[1:]
                    data = [r[1:] for r in data]
                width = max([len(columns), *(len(r) for r in data)])
                columns += [f"Column {i + 1}" for i in range(len(columns), width)]
                key = (row.get("equipment") or "", item.label, tuple(columns))
                bucket = groups.setdefault(key, [])
                for number, cells in enumerate(data, start=1):
                    bucket.append([row["booking_id"], set_index, number, *cells, *[""] * (width - len(cells))])
    ordered = sorted(groups.items(), key=lambda kv: (kv[0][0].lower(), kv[0][1].lower()))
    return [(f"{equipment} — {label}" if equipment else label, list(cols), data)
            for (equipment, label, cols), data in ordered]


def render_xlsx(columns, rows, *, summary, generated_at, status_counts=(), view_label="", charges=False) -> bytes:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    _register_xlsx_styles(wb)

    bookings = _SheetWriter(wb, "Bookings")
    bookings.header([c.label for c in columns])
    kinds = [c.kind if c.kind in ("int", "money", "datetime") else "text" for c in columns]
    for i, row in enumerate(rows):
        bookings.row([_xlsx_value(c, row.get(c.key)) for c in columns], kinds, striped=i % 2 == 1)
    bookings.ws.freeze_panes = "C2"
    bookings.flush(autofilter_columns=len(columns))

    labels = input_field_labels(rows)
    sets = _SheetWriter(wb, "Sample sets")
    sets.header(["Booking ID", "Equipment", "Sample set", *labels])
    for i, (row, number, values) in enumerate(_sample_set_rows(rows, labels)):
        sets.row([row["booking_id"], _xlsx_text(row.get("equipment") or ""), number, *values],
                 ["text", "text", "int", *["text"] * len(labels)], striped=i % 2 == 1)
    if not sets.data_rows:
        sets.blank()
        sets.line("None of the exported bookings has user inputs.", "exp_note")
    sets.ws.freeze_panes = "D2"
    sets.flush(autofilter_columns=3 + len(labels))

    tables = _SheetWriter(wb, "Input tables", max_width=40)
    groups = _input_table_groups(rows)
    for title, cols, data in groups:
        tables.line(_xlsx_text(title), "exp_band")
        tables.header(["Booking ID", "Sample set", "Row #", *cols])
        for i, cells in enumerate(data):
            tables.row([cells[0], cells[1], cells[2], *(_xlsx_text(c) for c in cells[3:])],
                       ["text", "int", "int", *["text"] * len(cols)], striped=i % 2 == 1)
        tables.blank()
    if not groups:
        tables.line("None of the exported bookings has a table input.", "exp_note")
    tables.flush()

    if charges:
        sheet = _SheetWriter(wb, "Charges", max_width=70)
        sheet.header(["Booking ID", "Equipment", "Line", "Description", "Amount (₹)"])
        stripe = False
        for row in rows:
            detail = row.get("detail")
            if detail is None or not detail.charges:
                continue
            for number, (description, amount) in enumerate(detail.charges, start=1):
                sheet.row(
                    [row["booking_id"], _xlsx_text(row.get("equipment") or ""), number, _xlsx_text(description),
                     float(amount) if amount is not None else None],
                    ["text", "text", "int", "text", "money"], striped=stripe,
                )
            stripe = not stripe
        if not sheet.data_rows:
            sheet.blank()
            sheet.line("None of the exported bookings has a charge breakdown.", "exp_note")
        sheet.ws.freeze_panes = "A2"
        sheet.flush(autofilter_columns=5)

    info = _SheetWriter(wb, "Filters", max_width=90)
    info.line(f"Bookings — {PORTAL_HEADER}", "exp_title")
    info.blank()
    details = [("Generated at (IST)", generated_at)]
    if view_label:
        details.append(("Exported from", view_label))
    details.append(("Rows", str(len(rows))))
    for label, value in [*details, *summary]:
        info.row([label, _xlsx_text(value)], ["text", "text"])
        info.rows[-1][0] = (label, "exp_label")
    if status_counts:
        info.blank()
        info.line("Bookings by status", "exp_band")
        for label, count in status_counts:
            info.row([label, count], ["text", "int"])
            info.rows[-1][0] = (label, "exp_label")
    info.blank()
    info.line("Sheets in this file", "exp_band")
    sheet_notes = [
        ("Bookings", "One row per booking: list columns, then one column per user input field (by equipment)."),
        ("Sample sets", "One row per sample set with that set's parameters."),
        ("Input tables", "Table inputs, one row per table row, grouped by equipment and field."),
    ]
    if charges:
        sheet_notes.append(("Charges", "Charge breakdown lines of each booking (₹)."))
    for label, text in sheet_notes:
        info.row([label, text], ["text", "text"])
        info.rows[-1][0] = (label, "exp_label")
    info.widths[0] = max(info.widths.get(0, 0), 22)
    info.flush()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def export_booking_list(request, queryset, *, view: str, export_format: str) -> HttpResponse:
    from .api_views import _user_is_accounts_finance_user

    params = request.query_params
    status_value = (params.get("status") or "").strip().upper()
    include_waitlist = (
        view == "my"
        and not _user_is_accounts_finance_user(request.user)
        and status_value in ("", BookingStatus.WAITLISTED)
    )
    entries = _waitlist_entries(request, include=include_waitlist)

    total = queryset.count() + len(entries)
    if total > EXPORT_ROW_LIMIT:
        raise ExportTooLarge(
            f"{total:,} bookings match these filters; exports are limited to "
            f"{EXPORT_ROW_LIMIT:,}. Narrow the filters (for example a date range) and try again.",
        )
    if export_format == "pdf" and total > PDF_ROW_LIMIT:
        raise ExportTooLarge(
            f"{total:,} bookings match these filters; the PDF has a details page for each booking and is "
            f"limited to {PDF_ROW_LIMIT:,}. Download Excel instead (up to {EXPORT_ROW_LIMIT:,} bookings) "
            "or narrow the filters (for example a date range).",
        )

    from .booking_export_details import build_booking_details

    show_charges = charges_visible(view, request.user)
    bookings = _load_bookings(queryset)
    rows = build_booking_rows(bookings)
    attach_details(rows, build_booking_details(bookings, include_charges=show_charges))
    if entries:
        search = (params.get("search") or "").strip()
        waitlist_rows = build_waitlist_rows(entries, search if len(search) >= 2 else "")
        ordering = (params.get("ordering") or "-created_at").strip()
        rows = rows + waitlist_rows
        # My Bookings interleaves waitlist entries by date only for the default created-at sorts.
        if ordering in ("created_at", "-created_at"):
            rows.sort(key=lambda r: r["booked_on"] or timezone.now(), reverse=ordering.startswith("-"))
    for index, row in enumerate(rows, start=1):
        row["sno"] = index

    columns = export_columns(view, request.user)
    now = timezone.now()
    generated_at = format_ist_datetime(now)
    summary = filters_summary(request, view=view)
    counts = status_counts(rows)
    view_label = "View Booking" if view == "staff" else "My Bookings"
    if export_format == "csv":
        extra = detail_columns(rows, charges=show_charges, include_charge_lines=True)
        content = render_csv(columns + extra, rows)
    elif export_format == "xlsx":
        extra = detail_columns(rows, charges=show_charges, include_charge_lines=False)
        content = render_xlsx(
            columns + extra, rows, summary=summary, generated_at=generated_at, status_counts=counts,
            view_label=view_label, charges=show_charges,
        )
    else:
        from .booking_export_pdf import render_pdf

        content = render_pdf(
            columns, rows, summary=summary, generated_at=generated_at, status_counts=counts,
            view_label=view_label, charges=show_charges,
        )

    response = HttpResponse(content, content_type=_CONTENT_TYPES[export_format])
    response["Content-Disposition"] = f'attachment; filename="{export_filename(export_format, now)}"'
    response["X-Export-Row-Count"] = str(len(rows))
    response["Cache-Control"] = "no-store"
    return response
